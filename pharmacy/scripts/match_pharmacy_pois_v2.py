from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
import sqlite3
import sys
import time
import unicodedata
from collections import Counter, defaultdict
from copy import copy
from datetime import datetime, timezone
from pathlib import Path

import geopandas as gpd
import osmium
import pandas as pd
import requests
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from rapidfuzz import fuzz, process
from shapely import wkb, wkt
from shapely.geometry import Point, shape


VERSION = "2.0.7"
PHARMACY_DIR = Path(__file__).resolve().parents[1]
OUTPUT_DIR = PHARMACY_DIR / "output"
CACHE_DIR = PHARMACY_DIR / "cache"
CHECKPOINT_DIR = CACHE_DIR / "checkpoints"
SOURCE_PATH = PHARMACY_DIR / "univero_pharmacy_crm.xlsx"
V1_PATH = OUTPUT_DIR / "PHARMACY_MASTER_GEOCODED.xlsx"
V1_CHECKPOINT = CHECKPOINT_DIR / "geocoded_pharmacies_cartociudad_v7.csv"
OUTPUT_PATH = OUTPUT_DIR / "PHARMACY_MASTER_GEOCODED_V2.xlsx"
TEMP_OUTPUT_PATH = OUTPUT_DIR / "PHARMACY_MASTER_GEOCODED_V2.tmp.xlsx"
LOG_PATH = OUTPUT_DIR / "geocoding_run_log_v2.txt"
CACHE_PATH = CACHE_DIR / "poi_match_cache_v2.sqlite"
OSM_PBF_PATH = CACHE_DIR / "osm" / "spain-latest.osm.pbf"
OSM_PART_PATH = CACHE_DIR / "osm" / "spain-latest.osm.pbf.part"
IGN_API = "https://api-features.ign.es/collections/administrativeunit"
CARTO_CANDIDATES = "https://www.cartociudad.es/geocoder/api/geocoder/candidates"
CARTO_FIND = "https://www.cartociudad.es/geocoder/api/geocoder/find"
GEOFABRIK_PBF = "https://download.geofabrik.de/europe/spain-latest.osm.pbf"

EXPECTED_TOTAL = 22356
EXPECTED_EXCEPTIONS = 8791
VALID_THRESHOLD = 0.95
WARNING_THRESHOLD = 0.90
MAX_RETRIES = 4
TIMEOUT = 20
REQUEST_DELAY = 1.0
IGN_PAGE_SIZE = 200
CHECKPOINT_BLOCK_SIZE = 100
CARTO_LIMIT = 10
OSM_ATTRIBUTION = "© OpenStreetMap contributors, ODbL"
IGN_ATTRIBUTION = "BDLJE CC-BY 4.0 ign.es"
CARTO_ATTRIBUTION = "CC BY 4.0 www.scne.es/productos.html#CartoCiudad"

SOURCE_TO_V1 = {
    "crm_id": "pharmacy_id",
    "account_name": "pharmacy_name",
    "sbrick_code": "sb_code_2026",
    "calle": "street_name_raw",
    "num": "street_number_raw",
    "postal_code": "postal_code_raw",
    "ciudad": "municipality_raw",
    "provincia": "province_raw",
    "ccaaa": "autonomous_region_raw",
}
RAW_SOURCE_COLUMNS = (
    "crm_id", "account_name", "sales_arena", "account_type", "iqvia_id",
    "sbrick_code", "ccaaa", "provincia", "ciudad", "postal_code",
    "calle", "num", "phone", "email",
)

NEW_COLUMNS = [
    "poi_match_provider", "poi_match_attempt", "poi_match_query",
    "poi_match_query_variant", "poi_match_status", "poi_match_quality",
    "poi_match_score", "poi_candidate_id", "poi_candidate_name",
    "poi_candidate_type", "poi_candidate_geometry_type", "poi_matched_address",
    "poi_matched_postal_code", "poi_matched_municipality",
    "poi_matched_province", "poi_name_match", "poi_municipality_match",
    "poi_province_match", "poi_postal_code_match", "poi_latitude",
    "poi_longitude", "poi_match_alternatives_json", "poi_match_notes",
    "poi_suggested_action", "poi_match_timestamp",
    "latitude_final", "longitude_final", "coordinate_source_final",
    "coordinate_quality_final",
]

PROVINCE_ALIASES = {
    "araba": "alava", "la coruna": "a coruna", "la coruna": "a coruna",
    "gerona": "girona", "lerida": "lleida", "orense": "ourense",
    "vizcaya": "bizkaia", "guipuzcoa": "gipuzkoa",
    "baleares": "illes balears", "islas baleares": "illes balears",
    "sta cruz tenerife": "santa cruz de tenerife",
    "santa cruz tenerife": "santa cruz de tenerife",
}

STATS = Counter()
_LOG_HANDLE = None
_LAST_CARTO_REQUEST = 0.0
_LAST_IGN_REQUEST = 0.0


def now_iso() -> str:
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")


def ensure_dirs() -> None:
    for path in (OUTPUT_DIR, CACHE_DIR, CHECKPOINT_DIR, OSM_PBF_PATH.parent):
        path.mkdir(parents=True, exist_ok=True)


def log(message: str) -> None:
    ensure_dirs()
    line = f"[{now_iso()}] {message}"
    print(line, flush=True)
    with LOG_PATH.open("a", encoding="utf-8") as handle:
        handle.write(line + "\n")


def clean_text(value) -> str:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return ""
    text = str(value).replace("\r", " ").replace("\n", " ").replace("\t", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return re.sub(r"\s*,\s*", ", ", text).strip(" ,")


def fold_text(value) -> str:
    text = clean_text(value).casefold()
    text = "".join(
        char for char in unicodedata.normalize("NFD", text)
        if unicodedata.category(char) != "Mn"
    )
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def normalize_province(value) -> str:
    text = fold_text(value)
    return PROVINCE_ALIASES.get(text, text)


def normalize_municipality(value) -> str:
    return fold_text(value)


def normalize_pharmacy_name(value) -> str:
    text = fold_text(value)
    while True:
        updated = re.sub(r"^(?:fcia|farmacia|pharmacy)\s+", "", text)
        if updated == text:
            break
        text = updated
    return text.strip()


def query_components(row) -> tuple[str, str, str, str]:
    name = clean_text(row.get("pharmacy_name", ""))
    variant = "SOURCE_NAME"
    prefix = re.match(r"^\s*FCIA\.\s*(.+)$", name, flags=re.IGNORECASE)
    if prefix:
        name = f"Farmacia {prefix.group(1).strip()}"
        variant = "FCIA_PREFIX_EXPANDED_TO_FARMACIA"
    municipality = clean_text(row.get("municipality_normalized", ""))
    province = clean_text(row.get("province_normalized", ""))
    return name, municipality, province, variant


def make_query_key(row) -> str:
    name, municipality, province, _ = query_components(row)
    parts = [unicodedata.normalize("NFC", value) for value in (name, municipality, province)]
    return hashlib.sha256(json.dumps(parts, ensure_ascii=False, separators=(",", ":")).encode("utf-8")).hexdigest()


def build_query(row) -> tuple[str, str]:
    name, municipality, province, variant = query_components(row)
    return f"{name}, {municipality}, {province}", variant


def normalized_query_text(query: str) -> str:
    return unicodedata.normalize("NFC", clean_text(query)).casefold()


def source_mapping() -> dict:
    return {
        "crm_id": "pharmacy_id",
        "account_name": "pharmacy_name",
        "sbrick_code": "sb_code_2026",
        "calle": "street_name_raw",
        "num": "street_number_raw",
        "postal_code": "postal_code_raw",
        "ciudad": "municipality_raw",
        "provincia": "province_raw",
        "ccaaa": "autonomous_region_raw",
    }


def read_v1_master() -> tuple[pd.DataFrame, str]:
    if V1_PATH.exists():
        try:
            book = pd.ExcelFile(V1_PATH)
            if "PHARMACY_MASTER" not in book.sheet_names:
                raise RuntimeError("V1 workbook lacks PHARMACY_MASTER sheet")
            df = pd.read_excel(
                V1_PATH, sheet_name="PHARMACY_MASTER", dtype=str,
                keep_default_na=False,
            )
            return df, str(V1_PATH)
        except Exception as exc:
            log(f"V1 workbook unavailable ({type(exc).__name__}); testing validated checkpoint fallback")
    if not V1_CHECKPOINT.exists():
        raise RuntimeError("STOPPER: neither readable V1 workbook nor V1 checkpoint exists")
    return pd.read_csv(V1_CHECKPOINT, dtype=str, keep_default_na=False), str(V1_CHECKPOINT)


def validate_inputs() -> tuple[pd.DataFrame, str, dict]:
    source = pd.read_excel(
        SOURCE_PATH, sheet_name="Sheet1", dtype=str, keep_default_na=False
    )
    v1, selected = read_v1_master()
    if len(source) != EXPECTED_TOTAL or len(v1) != EXPECTED_TOTAL:
        raise RuntimeError(
            f"STOPPER: row count differs from baseline; source={len(source)}, V1={len(v1)}"
        )
    if source["crm_id"].astype(str).str.strip().eq("").any():
        raise RuntimeError("STOPPER: blank CRM pharmacy IDs in source")
    if source["crm_id"].duplicated().any():
        raise RuntimeError("STOPPER: duplicate CRM pharmacy IDs in source")
    if "pharmacy_id" not in v1 or v1["pharmacy_id"].duplicated().any():
        raise RuntimeError("STOPPER: missing or duplicate pharmacy_id in V1 baseline")
    source_ids = set(source["crm_id"].astype(str))
    v1_ids = set(v1["pharmacy_id"].astype(str))
    if source_ids != v1_ids:
        raise RuntimeError("STOPPER: source and V1 pharmacy_id sets differ")

    source_index = source.set_index("crm_id").sort_index()
    v1_index = v1.set_index("pharmacy_id").sort_index()
    raw_mismatches = {}
    for column in RAW_SOURCE_COLUMNS:
        if column not in source.columns:
            raw_mismatches[column] = "missing from source"
            continue
        if column == "crm_id":
            count = int((source_index.index.astype(str) != v1_index.index.astype(str)).sum())
            if count:
                raw_mismatches[column] = count
            continue
        if column not in v1_index.columns:
            raw_mismatches[column] = "missing column"
            continue
        left = source_index[column].astype(str)
        right = v1_index[column].astype(str)
        count = int((left.to_numpy() != right.to_numpy()).sum())
        if count:
            raw_mismatches[column] = count
    for raw_column, standard_column in SOURCE_TO_V1.items():
        if standard_column not in v1.columns:
            raw_mismatches[standard_column] = "missing mapped column"
            continue
        if raw_column == "crm_id":
            if list(source_index.index.astype(str)) != list(v1_index.index.astype(str)):
                raw_mismatches[standard_column] = "ID mapping mismatch"
            continue
        if raw_column == "sbrick_code":
            left = source_index[raw_column].astype(str)
            right = v1_index[standard_column].astype(str)
            count = int((left.to_numpy() != right.to_numpy()).sum())
            if count:
                raw_mismatches[standard_column] = count
    if raw_mismatches:
        raise RuntimeError(f"STOPPER: V1/source raw reconciliation failed: {raw_mismatches}")

    required = (
        "pharmacy_id", "pharmacy_name", "sb_code_2026", "municipality_normalized",
        "province_normalized", "geocoding_quality", "latitude", "longitude",
    )
    missing = [column for column in required if column not in v1.columns]
    if missing:
        raise RuntimeError(f"STOPPER: V1 columns missing: {missing}")
    quality = v1["geocoding_quality"].astype(str).str.upper().str.strip()
    allowed = {"EXACT", "HIGH", "REVIEW", "FAILED"}
    if not set(quality.unique()).issubset(allowed):
        raise RuntimeError(f"STOPPER: unexpected V1 quality labels: {sorted(set(quality)-allowed)}")
    exception_mask = quality.isin({"REVIEW", "FAILED"})
    exception_count = int(exception_mask.sum())
    if exception_count != EXPECTED_EXCEPTIONS:
        raise RuntimeError(
            f"STOPPER: expected {EXPECTED_EXCEPTIONS} V1 exceptions, found {exception_count}"
        )
    exception_rows = v1.loc[exception_mask].copy()
    exception_rows["_query_key"] = exception_rows.apply(make_query_key, axis=1)
    unique_query_count = int(exception_rows["_query_key"].nunique())
    if unique_query_count != 8789:
        raise RuntimeError(f"STOPPER: expected 8789 unique name/locality keys, found {unique_query_count}")
    missing_locality = (
        exception_rows["municipality_normalized"].map(clean_text).eq("")
        | exception_rows["province_normalized"].map(clean_text).eq("")
    )
    mapping = source_mapping()
    log(f"Source selected: {SOURCE_PATH}; sheet=Sheet1; rows={len(source)}; columns={len(source.columns)}")
    log(f"V1 baseline selected: {selected}; rows={len(v1)}; IDs=unique {v1['pharmacy_id'].nunique()}")
    log(f"Column mapping: {json.dumps(mapping, ensure_ascii=False)}")
    log("Raw source-to-V1 reconciliation: all 14 original columns match exactly by ID; pharmacy_id and SB code mappings verified")
    log(
        f"V1 quality distribution: {quality.value_counts().to_dict()}; "
        f"exceptions={exception_count}; unique_query_keys={unique_query_count}; "
        f"insufficient_locality={int(missing_locality.sum())}"
    )
    log("External query fields: pharmacy_name + municipality_normalized + province_normalized; no CRM ID/contact/address fields; Spain is implicit in the locality context")
    return v1, selected, {"exceptions": exception_rows, "quality": quality}


def open_cache() -> sqlite3.Connection:
    ensure_dirs()
    conn = sqlite3.connect(CACHE_PATH)
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA synchronous=FULL")
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS api_cache (
            provider TEXT NOT NULL,
            query_key TEXT NOT NULL,
            query_text TEXT NOT NULL,
            status_code INTEGER NOT NULL,
            payload_json TEXT,
            error_text TEXT,
            attempts INTEGER NOT NULL DEFAULT 1,
            created_at TEXT NOT NULL,
            PRIMARY KEY (provider, query_key)
        );
        CREATE TABLE IF NOT EXISTS match_results (
            query_key TEXT PRIMARY KEY,
            matcher_version TEXT NOT NULL,
            result_json TEXT NOT NULL,
            updated_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS ign_pages (
            page_offset INTEGER PRIMARY KEY,
            feature_count INTEGER NOT NULL,
            saved_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS ign_municipalities (
            feature_id TEXT PRIMARY KEY,
            nationalcode TEXT,
            nameunit TEXT NOT NULL,
            codnut3 TEXT,
            geometry_wkb BLOB NOT NULL
        );
        CREATE TABLE IF NOT EXISTS osm_metadata (
            key TEXT PRIMARY KEY,
            value TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS osm_pharmacies (
            osm_type TEXT NOT NULL,
            osm_id TEXT NOT NULL,
            name TEXT NOT NULL,
            name_key TEXT NOT NULL,
            lat REAL,
            lon REAL,
            is_point INTEGER NOT NULL,
            tags_json TEXT NOT NULL,
            PRIMARY KEY (osm_type, osm_id)
        );
        CREATE INDEX IF NOT EXISTS idx_osm_name_key ON osm_pharmacies(name_key);
        """
    )
    conn.commit()
    return conn


def cached_api(conn: sqlite3.Connection, provider: str, query: str):
    key = hashlib.sha256(f"{provider}|{normalized_query_text(query)}".encode("utf-8")).hexdigest()
    row = conn.execute(
        "SELECT status_code,payload_json,error_text,attempts FROM api_cache WHERE provider=? AND query_key=?",
        (provider, key),
    ).fetchone()
    if row:
        STATS["cache_hits"] += 1
        payload = json.loads(row[1]) if row[1] else None
        return key, (row[0], payload, row[2], row[3])
    return key, None


def save_api_cache(conn, provider, key, query, status, payload, error, attempts):
    conn.execute(
        "INSERT OR REPLACE INTO api_cache(provider,query_key,query_text,status_code,payload_json,error_text,attempts,created_at) VALUES(?,?,?,?,?,?,?,?)",
        (provider, key, query, int(status), json.dumps(payload, ensure_ascii=False) if payload is not None else None,
         error, int(attempts), now_iso()),
    )
    conn.commit()


def rate_wait(provider: str) -> None:
    global _LAST_CARTO_REQUEST, _LAST_IGN_REQUEST
    delay = REQUEST_DELAY if provider.startswith("cartociudad") else 0.25
    last = _LAST_CARTO_REQUEST if provider.startswith("cartociudad") else _LAST_IGN_REQUEST
    remaining = delay - (time.monotonic() - last)
    if remaining > 0:
        time.sleep(remaining)
    if provider.startswith("cartociudad"):
        _LAST_CARTO_REQUEST = time.monotonic()
    else:
        _LAST_IGN_REQUEST = time.monotonic()


def parse_list_payload(data):
    if isinstance(data, list):
        return data
    if isinstance(data, dict):
        for key in ("response", "value", "candidates", "results", "result"):
            value = data.get(key)
            if isinstance(value, list):
                return value
    return []


def request_carto(conn, session, provider: str, url: str, query: str, limit: int = CARTO_LIMIT):
    query_key, cached = cached_api(conn, provider, query)
    if cached and cached[0] == 200 and cached[2] is None:
        return parse_list_payload(cached[1]), "", True
    params = {"q": query}
    if url == CARTO_CANDIDATES:
        params["limit"] = limit
    last_error = ""
    for attempt in range(1, MAX_RETRIES + 2):
        rate_wait("cartociudad")
        STATS["requests_sent"] += 1
        try:
            response = session.get(url, params=params, timeout=(10, TIMEOUT))
            if response.status_code == 429 or response.status_code >= 500:
                last_error = f"HTTP {response.status_code}"
                STATS["retries"] += int(attempt <= MAX_RETRIES)
                if response.status_code == 429:
                    STATS["http_429"] += 1
                else:
                    STATS["http_5xx"] += 1
                if attempt <= MAX_RETRIES:
                    time.sleep(min(2 ** (attempt - 1), 16) + REQUEST_DELAY)
                    continue
            response.raise_for_status()
            try:
                payload = response.json()
            except ValueError as exc:
                last_error = f"Invalid JSON: {exc}"
                payload = None
            if payload is not None:
                save_api_cache(conn, provider, query_key, query, response.status_code, payload, None, attempt)
                return parse_list_payload(payload), "", False
            save_api_cache(conn, provider, query_key, query, response.status_code, None, last_error, attempt)
            return [], last_error, False
        except requests.RequestException as exc:
            last_error = f"{type(exc).__name__}: {exc}"
            STATS["retries"] += int(attempt <= MAX_RETRIES)
            if attempt <= MAX_RETRIES:
                time.sleep(min(2 ** (attempt - 1), 16) + REQUEST_DELAY)
                continue
    STATS["errors"] += 1
    save_api_cache(conn, provider, query_key, query, 0, None, last_error, MAX_RETRIES + 1)
    return [], last_error, False


def download_osm_pbf(session) -> tuple[Path | None, dict]:
    ensure_dirs()
    try:
        head = session.head(GEOFABRIK_PBF, timeout=(15, 45), allow_redirects=True)
        head.raise_for_status()
        expected_size = int(head.headers.get("Content-Length", "0"))
        last_modified = head.headers.get("Last-Modified", "unknown")
    except Exception as exc:
        log(f"OSM PBF HEAD failed: {type(exc).__name__}: {exc}")
        return None, {"error": str(exc)}
    if OSM_PBF_PATH.exists() and expected_size and OSM_PBF_PATH.stat().st_size == expected_size:
        log(f"Reusing local OSM PBF: bytes={expected_size}; last_modified={last_modified}")
        return OSM_PBF_PATH, {"bytes": expected_size, "last_modified": last_modified, "reused": True}

    downloaded = OSM_PART_PATH.stat().st_size if OSM_PART_PATH.exists() else 0
    headers = {"Range": f"bytes={downloaded}-"} if downloaded else {}
    try:
        response = session.get(
            GEOFABRIK_PBF, headers=headers, stream=True, timeout=(20, 120), allow_redirects=True
        )
        if downloaded and response.status_code != 206:
            response.close()
            downloaded = 0
            response = session.get(
                GEOFABRIK_PBF, stream=True, timeout=(20, 120), allow_redirects=True
            )
        response.raise_for_status()
        mode = "ab" if downloaded and response.status_code == 206 else "wb"
        current = downloaded if mode == "ab" else 0
        with OSM_PART_PATH.open(mode) as handle:
            for chunk in response.iter_content(chunk_size=8 * 1024 * 1024):
                if not chunk:
                    continue
                handle.write(chunk)
                current += len(chunk)
                if current and current % (100 * 1024 * 1024) < len(chunk):
                    log(f"OSM PBF download progress: {current:,}/{expected_size:,} bytes")
        response.close()
        if expected_size and OSM_PART_PATH.stat().st_size != expected_size:
            raise IOError(
                f"OSM download size mismatch: {OSM_PART_PATH.stat().st_size} != {expected_size}"
            )
        OSM_PART_PATH.replace(OSM_PBF_PATH)
    except Exception as exc:
        log(f"OSM PBF download failed; partial file retained for resume: {type(exc).__name__}: {exc}")
        return None, {"error": str(exc), "partial_bytes": OSM_PART_PATH.stat().st_size if OSM_PART_PATH.exists() else 0}

    digest = hashlib.sha256()
    with OSM_PBF_PATH.open("rb") as handle:
        while True:
            block = handle.read(16 * 1024 * 1024)
            if not block:
                break
            digest.update(block)
    metadata = {
        "bytes": OSM_PBF_PATH.stat().st_size,
        "last_modified": last_modified,
        "sha256": digest.hexdigest(),
        "url": GEOFABRIK_PBF,
        "reused": False,
    }
    log(f"OSM PBF download complete: bytes={metadata['bytes']:,}; last_modified={last_modified}; sha256={metadata['sha256']}")
    return OSM_PBF_PATH, metadata


def request_ign(session, url, params):
    last_error = ""
    for attempt in range(1, MAX_RETRIES + 2):
        rate_wait("ign")
        try:
            response = session.get(url, params=params, timeout=(15, 90))
            if response.status_code == 429 or response.status_code >= 500:
                last_error = f"HTTP {response.status_code}"
                if attempt <= MAX_RETRIES:
                    time.sleep(min(2 ** (attempt - 1), 16) + 0.25)
                    continue
            response.raise_for_status()
            return response.json()
        except (requests.RequestException, ValueError) as exc:
            last_error = f"{type(exc).__name__}: {exc}"
            if attempt <= MAX_RETRIES:
                time.sleep(min(2 ** (attempt - 1), 16) + 0.25)
    raise RuntimeError(f"IGN API failed after retries: {last_error}")


def load_ign_boundaries(conn, session) -> tuple[gpd.GeoDataFrame, dict]:
    collection = request_ign(session, IGN_API, {"f": "json"})
    crs_values = collection.get("crs") or []
    default_crs = crs_values[0] if crs_values else collection.get("storageCRS", "EPSG:4326")
    crs_text = str(default_crs).upper()
    source_crs = "EPSG:4258" if "4258" in crs_text else "EPSG:4326"
    first = request_ign(
        session, f"{IGN_API}/items",
        {"limit": IGN_PAGE_SIZE, "offset": 0, "f": "json",
         "properties": "country,nationalcode,nameunit,codnut3,nationallevelname"},
    )
    total = int(first.get("numberMatched") or 0)
    if total < 8000:
        raise RuntimeError(f"STOPPER: IGN municipality feature count unexpected: {total}")
    expected_offsets = list(range(0, total, IGN_PAGE_SIZE))
    for page_no, offset in enumerate(expected_offsets, start=1):
        done = conn.execute("SELECT 1 FROM ign_pages WHERE page_offset=?", (offset,)).fetchone()
        if done:
            continue
        page = first if offset == 0 else request_ign(
            session, f"{IGN_API}/items",
            {"limit": IGN_PAGE_SIZE, "offset": offset, "f": "json",
             "properties": "country,nationalcode,nameunit,codnut3,nationallevelname"},
        )
        features = page.get("features", [])
        stored = 0
        for feature in features:
            props = feature.get("properties", {})
            geom = feature.get("geometry")
            if props.get("country") != "ES" or not geom:
                continue
            if clean_text(props.get("nationallevelname")).casefold() != "municipio":
                continue
            try:
                geometry = shape(geom)
                if geometry.is_empty:
                    continue
                if not geometry.is_valid:
                    geometry = geometry.buffer(0)
                feature_id = str(feature.get("id") or props.get("nationalcode") or "")
                if not feature_id:
                    continue
                conn.execute(
                    "INSERT OR REPLACE INTO ign_municipalities(feature_id,nationalcode,nameunit,codnut3,geometry_wkb) VALUES(?,?,?,?,?)",
                    (feature_id, str(props.get("nationalcode") or ""),
                     clean_text(props.get("nameunit")), clean_text(props.get("codnut3")),
                     sqlite3.Binary(geometry.wkb)),
                )
                stored += 1
            except Exception as exc:
                STATS["ign_invalid_features"] += 1
                log(f"IGN feature skipped during geometry parsing: {type(exc).__name__}")
        conn.execute(
            "INSERT OR REPLACE INTO ign_pages(page_offset,feature_count,saved_at) VALUES(?,?,?)",
            (offset, stored, now_iso()),
        )
        conn.commit()
        if page_no == 1 or page_no % 5 == 0 or page_no == len(expected_offsets):
            log(f"IGN boundary pages: {page_no}/{len(expected_offsets)}; current page features={stored}")
    records = conn.execute(
        "SELECT feature_id,nationalcode,nameunit,codnut3,geometry_wkb FROM ign_municipalities"
    ).fetchall()
    if len(records) < 8000:
        raise RuntimeError(f"STOPPER: cached IGN municipal polygons incomplete: {len(records)}")
    frame = pd.DataFrame(
        [{"feature_id": r[0], "nationalcode": r[1], "nameunit": r[2], "codnut3": r[3]}
         for r in records]
    )
    geometry = [wkb.loads(bytes(r[4])) for r in records]
    boundaries = gpd.GeoDataFrame(frame, geometry=geometry, crs=source_crs)
    if source_crs != "EPSG:4326":
        boundaries = boundaries.to_crs("EPSG:4326")
    boundaries = boundaries[boundaries.geometry.notna() & ~boundaries.geometry.is_empty].reset_index(drop=True)
    boundaries.sindex
    metadata = {
        "source": IGN_API,
        "source_crs": source_crs,
        "feature_count": len(boundaries),
        "retrieved_at": now_iso(),
        "attribution": IGN_ATTRIBUTION,
    }
    log(f"IGN administrative polygons ready: features={len(boundaries)}; CRS=EPSG:4326; source={IGN_API}")
    return boundaries, metadata


def build_province_nuts_map(v1: pd.DataFrame, boundaries: gpd.GeoDataFrame) -> tuple[dict, dict]:
    by_name = defaultdict(set)
    for row in boundaries[["nameunit", "nationalcode"]].itertuples(index=False):
        national_code = re.sub(r"\D", "", clean_text(row.nationalcode))
        province_code = national_code[4:6] if len(national_code) == 11 else ""
        if province_code:
            by_name[normalize_municipality(row.nameunit)].add(province_code)
    support = defaultdict(Counter)
    direct_single_admin_keys = set()
    pairs = v1[["municipality_normalized", "province_normalized"]].drop_duplicates()
    for municipality, province in pairs.itertuples(index=False, name=None):
        muni_key = normalize_municipality(municipality)
        prov_key = normalize_province(province)
        codes = by_name.get(muni_key, set())
        if muni_key and prov_key and len(codes) == 1:
            code = next(iter(codes))
            support[prov_key][code] += 1
            if muni_key == prov_key and (prov_key, code) in {
                ("ceuta", "51"), ("melilla", "52")
            }:
                direct_single_admin_keys.add(prov_key)
    mapping = {}
    conflicts = {}
    summary = {}
    for province, counts in support.items():
        total = sum(counts.values())
        code, votes = counts.most_common(1)[0]
        share = votes / total if total else 0.0
        summary[province] = total
        direct_single_municipality_province = (
            total == 1
            and province in direct_single_admin_keys
            and code == ("51" if province == "ceuta" else "52")
        )
        if direct_single_municipality_province or (total >= 10 and share >= 0.85):
            mapping[province] = code
        else:
            conflicts[province] = {
                "municipality_pairs": total,
                "dominant_ine_province_code": code,
                "dominant_share": round(share, 4),
                "code_counts": dict(counts),
            }
    log(
        f"IGN province crosswalk: resolved={len(mapping)}; conflicting={len(conflicts)}; "
        "acceptance rule=at least 10 distinct municipality/province pairs and >=85% dominant code, "
        "or direct unique Ceuta/Melilla municipality"
    )
    if conflicts:
        log(f"IGN province crosswalk conflicts: {json.dumps(conflicts, ensure_ascii=False)}")
    return mapping, summary


def locate_point(lat, lon, boundaries: gpd.GeoDataFrame):
    try:
        lat = float(lat)
        lon = float(lon)
    except (TypeError, ValueError):
        return None, "INVALID_COORDINATES"
    if not math.isfinite(lat) or not math.isfinite(lon) or (lat == 0 and lon == 0):
        return None, "INVALID_COORDINATES"
    point = Point(lon, lat)
    try:
        indexes = boundaries.sindex.query(point)
    except Exception:
        indexes = []
    hits = []
    for index in indexes:
        feature = boundaries.iloc[int(index)]
        try:
            if feature.geometry.covers(point):
                hits.append(feature)
        except Exception:
            continue
    if not hits:
        return None, "OUTSIDE_IGN_MUNICIPAL_POLYGONS"
    if len(hits) > 1:
        hits.sort(key=lambda item: item.geometry.area)
    return hits[0], "INSIDE_ES"


def osm_name_key(value) -> str:
    return normalize_pharmacy_name(value)


def index_osm_pharmacies(conn: sqlite3.Connection, pbf_path: Path | None, osm_meta: dict) -> list[dict]:
    if not pbf_path or not pbf_path.exists():
        log("OSM local fallback unavailable: Spain PBF missing")
        return []
    sha = osm_meta.get("sha256")
    if not sha:
        digest = hashlib.sha256()
        with pbf_path.open("rb") as handle:
            for block in iter(lambda: handle.read(16 * 1024 * 1024), b""):
                digest.update(block)
        sha = digest.hexdigest()
    osm_meta["sha256"] = sha
    cached_hash = conn.execute("SELECT value FROM osm_metadata WHERE key='sha256'").fetchone()
    cached_count = conn.execute("SELECT COUNT(*) FROM osm_pharmacies").fetchone()[0]
    if cached_hash and cached_hash[0] == sha and cached_count:
        log(f"Reusing local OSM pharmacy index: features={cached_count}; sha256={sha}")
    else:
        conn.execute("DELETE FROM osm_pharmacies")
        conn.execute("DELETE FROM osm_metadata")
        conn.commit()
        batch = []
        total = 0

        def flush_batch():
            if batch:
                conn.executemany(
                    "INSERT OR REPLACE INTO osm_pharmacies(osm_type,osm_id,name,name_key,lat,lon,is_point,tags_json) VALUES(?,?,?,?,?,?,?,?)",
                    batch,
                )
                conn.commit()
                batch.clear()

        log(f"Parsing local OSM PBF for amenity=pharmacy; bytes={pbf_path.stat().st_size:,}")
        processor = osmium.FileProcessor(str(pbf_path)).with_filter(
            osmium.filter.TagFilter(("amenity", "pharmacy"))
        )
        type_names = {"n": "node", "w": "way", "r": "relation"}
        for obj in processor:
            tags = {tag.k: tag.v for tag in obj.tags}
            name = clean_text(tags.get("name") or tags.get("name:es") or tags.get("name:ca") or "")
            if not name:
                continue
            object_type = obj.type_str()
            is_point = object_type == "n"
            lat = lon = None
            if is_point:
                try:
                    if not obj.location.valid():
                        continue
                    lat, lon = float(obj.location.lat), float(obj.location.lon)
                except Exception:
                    continue
            keep_tags = {
                key: tags[key] for key in (
                    "amenity", "name", "name:es", "name:ca", "alt_name", "brand",
                    "operator", "addr:city", "addr:province", "addr:postcode", "addr:street",
                ) if tags.get(key)
            }
            batch.append((type_names.get(object_type, object_type), str(obj.id), name,
                          osm_name_key(name), lat, lon, int(is_point),
                          json.dumps(keep_tags, ensure_ascii=False)))
            total += 1
            if len(batch) >= 1000:
                flush_batch()
            if total and total % 10000 == 0:
                log(f"OSM tagged pharmacy entities indexed: {total:,}")
        flush_batch()
        conn.execute("INSERT OR REPLACE INTO osm_metadata(key,value) VALUES('sha256',?)", (sha,))
        conn.commit()
        log(f"OSM pharmacy extraction complete: named_features={total}; sha256={sha}")

    rows = conn.execute(
        "SELECT osm_type,osm_id,name,name_key,lat,lon,is_point,tags_json FROM osm_pharmacies"
    ).fetchall()
    features = []
    for row in rows:
        try:
            tags = json.loads(row[7])
        except Exception:
            tags = {}
        features.append({
            "osm_type": row[0], "osm_id": row[1], "name": row[2], "name_key": row[3],
            "lat": row[4], "lon": row[5], "is_point": bool(row[6]), "tags": tags,
        })
    return features


def candidate_name(candidate: dict) -> str:
    for key in ("name", "poiName", "namePoi", "nombre", "label", "title", "address"):
        value = clean_text(candidate.get(key))
        if value:
            return value
    return ""


def candidate_type(candidate: dict) -> str:
    values = []
    for key in ("type", "category", "class", "kind", "subtype", "typeName", "description"):
        value = clean_text(candidate.get(key))
        if value:
            values.append(value)
    return " | ".join(values)


def explicit_pharmacy(candidate: dict, name: str, type_text: str) -> bool:
    haystack = fold_text(" ".join((type_text, name)))
    return (
        "farmacia" in haystack or "pharmacy" in haystack
        or clean_text(candidate.get("amenity")).casefold() == "pharmacy"
    )


def candidate_coords(candidate: dict):
    try:
        lat = float(candidate.get("lat", candidate.get("latitude")))
        lon = float(candidate.get("lng", candidate.get("lon", candidate.get("longitude"))))
        return lat, lon
    except (TypeError, ValueError):
        return None, None


def geometry_type(candidate: dict) -> str:
    geometry = candidate.get("geom") or candidate.get("geometry")
    if isinstance(geometry, dict):
        return clean_text(geometry.get("type")).upper() or "UNKNOWN"
    if isinstance(geometry, str) and geometry.strip():
        try:
            return wkt.loads(geometry).geom_type.upper()
        except Exception:
            match = re.match(r"\s*([A-Za-z]+)", geometry)
            return match.group(1).upper() if match else "UNKNOWN"
    if candidate_coords(candidate)[0] is not None:
        if fold_text(candidate.get("type")) in {"portal", "point", "poi"}:
            return "POINT"
    return "UNKNOWN"


def admin_match(input_value, result_value, province=False) -> str:
    if not clean_text(input_value) or not clean_text(result_value):
        return "UNKNOWN"
    left = normalize_province(input_value) if province else normalize_municipality(input_value)
    right = normalize_province(result_value) if province else normalize_municipality(result_value)
    return "MATCH" if left == right else "MISMATCH"


def precision_decimals(value) -> int:
    text = str(value)
    if "e" in text.lower():
        return 0
    return len(text.partition(".")[2].rstrip("0")) if "." in text else 0


def score_candidate(row, candidate, provider, boundaries, province_map):
    name = candidate_name(candidate)
    type_text = candidate_type(candidate)
    name_key = normalize_pharmacy_name(name)
    source_key = normalize_pharmacy_name(row.get("pharmacy_name", ""))
    name_ratio = fuzz.ratio(source_key, name_key) if source_key and name_key else 0
    exact_name = bool(source_key and source_key == name_key)
    pharmacy_type = (
        clean_text(candidate.get("amenity")).casefold() == "pharmacy"
        if provider == "OpenStreetMap local"
        else explicit_pharmacy(candidate, name, type_text)
    )
    lat, lon = candidate_coords(candidate)
    geom_type = "POINT" if provider == "OpenStreetMap local" and candidate.get("is_point") else geometry_type(candidate)
    point_ok = bool(lat is not None and lon is not None and geom_type == "POINT")
    point_feature, territory = locate_point(lat, lon, boundaries) if point_ok else (None, "NO_POINT_GEOMETRY")
    inside_spain = point_feature is not None and territory == "INSIDE_ES"
    source_muni = clean_text(row.get("municipality_normalized", ""))
    source_prov = clean_text(row.get("province_normalized", ""))
    polygon_muni = clean_text(point_feature.get("nameunit")) if point_feature is not None else ""
    polygon_national_code = re.sub(
        r"\D", "", clean_text(point_feature.get("nationalcode"))
    ) if point_feature is not None else ""
    polygon_province_code = (
        polygon_national_code[4:6] if len(polygon_national_code) == 11 else ""
    )
    geometry_muni_match = admin_match(source_muni, polygon_muni) if polygon_muni else "UNKNOWN"

    if provider == "OpenStreetMap local":
        tags = candidate.get("tags", {})
        tagged_muni = tags.get("addr:city", "")
        tagged_prov = tags.get("addr:province", "")
        municipality_match = admin_match(source_muni, tagged_muni) if tagged_muni else geometry_muni_match
        mapped_prov_code = province_map.get(normalize_province(source_prov), "")
        province_match = admin_match(source_prov, tagged_prov, province=True) if tagged_prov else (
            "MATCH" if mapped_prov_code and polygon_province_code == mapped_prov_code else "UNKNOWN"
        )
        postal = clean_text(tags.get("addr:postcode", ""))
        address = " ".join(x for x in (tags.get("addr:street", ""), tags.get("addr:city", "")) if x)
        candidate_id = f"{candidate.get('osm_type')}/{candidate.get('osm_id')}"
        candidate_geometry = "POINT" if candidate.get("is_point") else "AREA_OR_RELATION"
    else:
        tags = {}
        returned_muni = clean_text(candidate.get("muni", candidate.get("municipality", "")))
        returned_prov = clean_text(candidate.get("province", ""))
        city_api_match = admin_match(source_muni, returned_muni) if returned_muni else "UNKNOWN"
        prov_api_match = admin_match(source_prov, returned_prov, province=True) if returned_prov else "UNKNOWN"
        municipality_match = geometry_muni_match
        if city_api_match == "MISMATCH":
            municipality_match = "MISMATCH"
        elif city_api_match == "MATCH" and geometry_muni_match in {"MATCH", "UNKNOWN"}:
            municipality_match = "MATCH"
        mapped_prov_code = province_map.get(normalize_province(source_prov), "")
        polygon_prov_match = "MATCH" if mapped_prov_code and polygon_province_code == mapped_prov_code else "UNKNOWN"
        province_match = prov_api_match if returned_prov else polygon_prov_match
        if prov_api_match == "MISMATCH":
            province_match = "MISMATCH"
        postal = clean_text(candidate.get("postalCode", candidate.get("postal_code", "")))
        address = clean_text(candidate.get("address", ""))
        candidate_id = clean_text(candidate.get("id", ""))
        candidate_geometry = geom_type

    source_postal = re.sub(r"\D", "", clean_text(row.get("postal_code_normalized", "")))
    result_postal = re.sub(r"\D", "", postal)
    postal_match = "UNKNOWN"
    if source_postal and result_postal:
        postal_match = "MATCH" if source_postal == result_postal else "MISMATCH"

    name_points = 40 if exact_name else (36 if name_ratio >= 95 else int(40 * name_ratio / 100))
    score = name_points + (20 if pharmacy_type else 0)
    score += 20 if municipality_match == "MATCH" else 0
    score += 10 if province_match == "MATCH" else 0
    score += 10 if inside_spain and point_ok else 0
    near_municipality_centroid = False
    if point_feature is not None:
        try:
            near_municipality_centroid = Point(lon, lat).distance(point_feature.geometry.centroid) <= 0.00004
        except Exception:
            pass

    hard_mismatch = municipality_match == "MISMATCH" or province_match == "MISMATCH"
    locality_confirmed = municipality_match == "MATCH" and province_match == "MATCH"
    if hard_mismatch or not inside_spain or not point_ok:
        quality = "FAILED" if hard_mismatch or territory == "OUTSIDE_IGN_MUNICIPAL_POLYGONS" or not point_ok else "REVIEW"
    elif not pharmacy_type or not locality_confirmed or postal_match == "MISMATCH" or near_municipality_centroid:
        quality = "REVIEW"
    elif name_ratio >= 100:
        quality = "EXACT"
    elif name_ratio >= 95:
        quality = "HIGH"
    else:
        quality = "REVIEW" if score >= 60 else "FAILED"

    return {
        "provider": provider,
        "candidate_id": candidate_id,
        "candidate_name": name,
        "candidate_type": type_text or ("amenity=pharmacy" if pharmacy_type else ""),
        "geometry_type": candidate_geometry,
        "address": address,
        "postal_code": postal,
        "municipality": polygon_muni or clean_text(candidate.get("muni", "")),
        "province": clean_text(candidate.get("province", "")) if provider != "OpenStreetMap local" else clean_text(tags.get("addr:province", "")),
        "name_match": "MATCH" if exact_name else ("HIGH_SIMILARITY" if name_ratio >= 95 else "MISMATCH"),
        "name_similarity": round(name_ratio, 2),
        "municipality_match": municipality_match,
        "province_match": province_match,
        "postal_match": postal_match,
        "latitude": float(lat) if lat is not None else None,
        "longitude": float(lon) if lon is not None else None,
        "score": min(100, int(score)),
        "quality": quality,
        "inside_spain": bool(inside_spain),
        "point_geometry": bool(point_ok),
        "near_municipality_centroid": bool(near_municipality_centroid),
        "source_postal": source_postal,
        "returned_postal": result_postal,
        "notes": "; ".join(
            reason for reason, condition in (
                ("candidate is not explicitly identified as a pharmacy", not pharmacy_type),
                ("municipality not confirmed by official polygon", municipality_match != "MATCH"),
                ("province not confirmed", province_match != "MATCH"),
                ("candidate outside IGN municipal polygons", territory == "OUTSIDE_IGN_MUNICIPAL_POLYGONS"),
                ("candidate lacks point geometry", not point_ok),
                ("postal code mismatch", postal_match == "MISMATCH"),
                ("possible municipal centroid", near_municipality_centroid),
                ("municipality mismatch", municipality_match == "MISMATCH"),
                ("province mismatch", province_match == "MISMATCH"),
            ) if condition
        ),
    }


def resolve_carto_nonpoint(row, candidates, conn, session):
    resolved = list(candidates)
    source_key = normalize_pharmacy_name(row.get("pharmacy_name", ""))
    for candidate in candidates:
        if not isinstance(candidate, dict):
            continue
        name = candidate_name(candidate)
        type_text = candidate_type(candidate)
        if not explicit_pharmacy(candidate, name, type_text):
            continue
        if candidate_coords(candidate)[0] is not None and geometry_type(candidate) == "POINT":
            continue
        if not name or fuzz.ratio(source_key, normalize_pharmacy_name(name)) < 90:
            continue
        municipality = clean_text(row.get("municipality_normalized", ""))
        province = clean_text(row.get("province_normalized", ""))
        find_query = f"{name}, {municipality}, {province}"
        found, error, cached = request_carto(
            conn, session, "cartociudad_find", CARTO_FIND, find_query
        )
        STATS["find_queries"] += 1
        if cached:
            STATS["find_cache_hits"] += 1
        if error:
            STATS["find_errors"] += 1
            continue
        for item in found:
            if isinstance(item, dict):
                merged = dict(candidate)
                merged.update(item)
                resolved.append(merged)
    return resolved


def rank_carto_candidates(row, candidates, boundaries, province_map, conn, session):
    candidates = resolve_carto_nonpoint(row, candidates, conn, session)
    evaluated = []
    for candidate in candidates:
        if not isinstance(candidate, dict):
            continue
        item = score_candidate(row, candidate, "CartoCiudad", boundaries, province_map)
        item["raw_candidate"] = candidate
        evaluated.append(item)
    eligible = [
        item for item in evaluated
        if item["quality"] in {"EXACT", "HIGH"}
        and item["name_similarity"] >= 95
        and item["point_geometry"]
        and item["inside_spain"]
    ]
    if len(eligible) == 1:
        strong = eligible[0]
        conflicting = [
            item for item in evaluated
            if item is not strong and item["name_similarity"] >= 95
            and item["candidate_id"] != strong["candidate_id"]
        ]
        if not conflicting:
            return strong, evaluated, ""
    if len(eligible) > 1:
        return None, evaluated, "multiple high-confidence CartoCiudad POIs"
    if evaluated:
        close = [item for item in evaluated if item["name_similarity"] >= 90]
        if close:
            return None, evaluated, "CartoCiudad returned near matches but none passed all hard checks"
    return None, evaluated, "no qualifying CartoCiudad candidate"


def match_osm(row, features, boundaries, province_map):
    source_key = normalize_pharmacy_name(row.get("pharmacy_name", ""))
    if not source_key:
        return None, [], "blank pharmacy name"
    exact = [feature for feature in features if feature["is_point"] and feature["name_key"] == source_key]
    if not exact:
        names = sorted({feature["name_key"] for feature in features if feature["is_point"] and feature["name_key"]})
        close = process.extract(source_key, names, scorer=fuzz.ratio, limit=15, score_cutoff=95)
        keys = {key for key, _, _ in close}
        exact = [feature for feature in features if feature["is_point"] and feature["name_key"] in keys]
    evaluated = []
    for feature in exact:
        candidate = {
            "name": feature["name"], "amenity": "pharmacy", "type": "amenity=pharmacy",
            "lat": feature["lat"], "lon": feature["lon"], "is_point": feature["is_point"],
            "osm_type": feature["osm_type"], "osm_id": feature["osm_id"], "tags": feature["tags"],
        }
        evaluated.append(score_candidate(row, candidate, "OpenStreetMap local", boundaries, province_map))
    eligible = [
        item for item in evaluated
        if item["quality"] in {"EXACT", "HIGH"} and item["name_similarity"] >= 95
        and item["inside_spain"] and item["point_geometry"]
    ]
    if len(eligible) == 1:
        return eligible[0], evaluated, ""
    if len(eligible) > 1:
        return None, evaluated, "multiple high-confidence OSM pharmacy nodes"
    area_matches = [
        feature for feature in features
        if not feature["is_point"] and feature["name_key"] == source_key
    ]
    if area_matches:
        return None, evaluated, "OSM name match is mapped as an area/relation; no point coordinate accepted"
    if evaluated:
        return None, evaluated, "OSM name candidates did not pass locality/geometry checks"
    return None, [], "no OSM pharmacy-name candidate"


def candidate_summary(item: dict | None) -> dict:
    if not item:
        return {}
    return {key: value for key, value in item.items() if key != "raw_candidate"}


def get_saved_match(conn, query_key):
    row = conn.execute(
        "SELECT matcher_version,result_json FROM match_results WHERE query_key=?", (query_key,)
    ).fetchone()
    if row and row[0] == VERSION:
        STATS["match_checkpoint_hits"] += 1
        return json.loads(row[1])
    return None


def save_match(conn, query_key, result):
    conn.execute(
        "INSERT OR REPLACE INTO match_results(query_key,matcher_version,result_json,updated_at) VALUES(?,?,?,?)",
        (query_key, VERSION, json.dumps(result, ensure_ascii=False), now_iso()),
    )


def match_one(row, conn, session, boundaries, province_map, osm_features, persist=True):
    key = make_query_key(row)
    if persist:
        saved = get_saved_match(conn, key)
        if saved is not None:
            return saved
    name = clean_text(row.get("pharmacy_name", ""))
    municipality = clean_text(row.get("municipality_normalized", ""))
    province = clean_text(row.get("province_normalized", ""))
    if not name or not municipality or not province:
        result = {
            "status": "SKIPPED_INCOMPLETE", "quality": "REVIEW", "score": 0,
            "provider": "", "query": "", "query_variant": "", "candidate": {},
            "alternatives": [], "notes": "pharmacy name, municipality or province missing; no query sent",
        }
        if persist:
            save_match(conn, key, result)
            conn.commit()
        return result

    query, variant = build_query(row)
    candidates, error, cached = request_carto(
        conn, session, "cartociudad_candidates", CARTO_CANDIDATES, query, CARTO_LIMIT
    )
    STATS["carto_query_count"] += 1
    if cached:
        STATS["carto_query_cache_hits"] += 1
    selected, evaluated, carto_note = rank_carto_candidates(
        row, candidates, boundaries, province_map, conn, session
    )
    for item in evaluated:
        if item["latitude"] is not None:
            STATS["candidate_coordinate_count"] += 1
        if item["municipality_match"] == "MISMATCH":
            STATS["municipality_mismatches"] += 1
        if item["province_match"] == "MISMATCH":
            STATS["province_mismatches"] += 1
        if item["postal_match"] == "MISMATCH":
            STATS["postcode_mismatches"] += 1
        if item["latitude"] is not None and not item["inside_spain"]:
            STATS["outside_candidate_coords"] += 1

    carto_close = [item for item in evaluated if item["name_similarity"] >= 90]
    if selected:
        result = {
            "status": "MATCHED", "quality": selected["quality"], "score": selected["score"],
            "provider": "CartoCiudad", "query": query, "query_variant": variant,
            "candidate": candidate_summary(selected),
            "alternatives": [candidate_summary(item) for item in evaluated],
            "notes": selected["notes"] or "unique pharmacy POI and administrative locality verified",
        }
        STATS["carto_accepted"] += 1
    else:
        osm_selected, osm_evaluated, osm_note = match_osm(row, osm_features, boundaries, province_map)
        for item in osm_evaluated:
            if item["latitude"] is not None:
                STATS["candidate_coordinate_count"] += 1
            if item["municipality_match"] == "MISMATCH":
                STATS["municipality_mismatches"] += 1
            if item["province_match"] == "MISMATCH":
                STATS["province_mismatches"] += 1
            if item["postal_match"] == "MISMATCH":
                STATS["postcode_mismatches"] += 1
            if item["latitude"] is not None and not item["inside_spain"]:
                STATS["outside_candidate_coords"] += 1
        conflict = bool(carto_close and osm_selected)
        if conflict:
            result = {
                "status": "AMBIGUOUS", "quality": "REVIEW", "score": max(x["score"] for x in carto_close),
                "provider": "CartoCiudad + OpenStreetMap local", "query": query,
                "query_variant": variant, "candidate": {},
                "alternatives": [candidate_summary(item) for item in carto_close + osm_evaluated],
                "notes": "CartoCiudad and OSM candidates conflict; no source chosen automatically",
            }
        elif osm_selected:
            result = {
                "status": "MATCHED", "quality": osm_selected["quality"], "score": osm_selected["score"],
                "provider": "OpenStreetMap local", "query": query, "query_variant": variant,
                "candidate": candidate_summary(osm_selected),
                "alternatives": [candidate_summary(item) for item in osm_evaluated],
                "notes": osm_selected["notes"] or "unique local pharmacy node; IGN municipality/province verified",
            }
            STATS["osm_accepted"] += 1
        elif error:
            quality = "REVIEW" if osm_evaluated else "FAILED"
            result = {
                "status": "API_ERROR", "quality": quality, "score": 0,
                "provider": "CartoCiudad + OpenStreetMap local", "query": query,
                "query_variant": variant, "candidate": {},
                "alternatives": [candidate_summary(item) for item in osm_evaluated],
                "notes": f"CartoCiudad error ({error}); OSM: {osm_note}",
            }
        else:
            near_candidates = carto_close + osm_evaluated
            if near_candidates or "area/relation" in osm_note:
                best = max(near_candidates, key=lambda item: item["score"], default={})
                result = {
                    "status": "AMBIGUOUS" if "multiple" in (carto_note + osm_note) else "REVIEW",
                    "quality": "REVIEW", "score": best.get("score", 0),
                    "provider": "CartoCiudad + OpenStreetMap local", "query": query,
                    "query_variant": variant, "candidate": candidate_summary(best),
                    "alternatives": [candidate_summary(item) for item in near_candidates],
                    "notes": f"{carto_note}; {osm_note}",
                }
            else:
                result = {
                    "status": "NO_RESULT", "quality": "FAILED", "score": 0,
                    "provider": "CartoCiudad + OpenStreetMap local", "query": query,
                    "query_variant": variant, "candidate": {}, "alternatives": [],
                    "notes": f"{carto_note}; {osm_note}",
                }
    result["timestamp"] = now_iso()
    STATS[f"quality_{result['quality'].lower()}"] += 1
    if persist:
        save_match(conn, key, result)
        conn.commit()
    return result


def pilot_keys(exception_rows: pd.DataFrame, pilot_size: int) -> list[str]:
    unique = exception_rows.drop_duplicates("_query_key").copy()
    if pilot_size >= len(unique):
        return unique["_query_key"].tolist()
    unique["_province_sort"] = unique["province_normalized"].map(normalize_province)
    unique = unique.sort_values(["_province_sort", "_query_key"])
    first_by_province = unique.drop_duplicates("_province_sort")
    chosen = first_by_province.head(pilot_size)
    if len(chosen) < pilot_size:
        remaining = unique[~unique["_query_key"].isin(chosen["_query_key"])]
        chosen = pd.concat([chosen, remaining.head(pilot_size - len(chosen))])
    return chosen["_query_key"].head(pilot_size).tolist()


def checkpoint_results(conn, key_rows: dict, completed_keys: list[str], pilot=False):
    if not completed_keys:
        return
    records = []
    for key in completed_keys:
        row = conn.execute("SELECT result_json FROM match_results WHERE query_key=?", (key,)).fetchone()
        if not row:
            continue
        result = json.loads(row[0])
        for source_row in key_rows[key]:
            records.append({
                "pharmacy_id": clean_text(source_row.get("pharmacy_id")),
                "query_key": key,
                "poi_match_status": result.get("status", ""),
                "poi_match_quality": result.get("quality", ""),
                "poi_match_score": result.get("score", 0),
                "poi_match_provider": result.get("provider", ""),
                "poi_candidate_name": result.get("candidate", {}).get("candidate_name", ""),
                "latitude": result.get("candidate", {}).get("latitude", ""),
                "longitude": result.get("candidate", {}).get("longitude", ""),
                "notes": result.get("notes", ""),
            })
    path = CHECKPOINT_DIR / ("poi_match_pilot_v2.csv" if pilot else "poi_match_checkpoint_v2.csv")
    if records:
        pd.DataFrame(records).to_csv(path, index=False, encoding="utf-8-sig")


def result_to_columns(result: dict) -> dict:
    candidate = result.get("candidate") or {}
    return {
        "poi_match_provider": result.get("provider", ""),
        "poi_match_attempt": result.get("status", ""),
        "poi_match_query": result.get("query", ""),
        "poi_match_query_variant": result.get("query_variant", ""),
        "poi_match_status": result.get("status", ""),
        "poi_match_quality": result.get("quality", ""),
        "poi_match_score": result.get("score", 0),
        "poi_candidate_id": candidate.get("candidate_id", ""),
        "poi_candidate_name": candidate.get("candidate_name", ""),
        "poi_candidate_type": candidate.get("candidate_type", ""),
        "poi_candidate_geometry_type": candidate.get("geometry_type", ""),
        "poi_matched_address": candidate.get("address", ""),
        "poi_matched_postal_code": candidate.get("postal_code", ""),
        "poi_matched_municipality": candidate.get("municipality", ""),
        "poi_matched_province": candidate.get("province", ""),
        "poi_name_match": candidate.get("name_match", ""),
        "poi_municipality_match": candidate.get("municipality_match", ""),
        "poi_province_match": candidate.get("province_match", ""),
        "poi_postal_code_match": candidate.get("postal_match", ""),
        "poi_latitude": candidate.get("latitude", ""),
        "poi_longitude": candidate.get("longitude", ""),
        "poi_match_alternatives_json": json.dumps(result.get("alternatives", []), ensure_ascii=False),
        "poi_match_notes": result.get("notes", ""),
        "poi_suggested_action": "" if result.get("quality") in {"EXACT", "HIGH"} else "Manual identity, locality and point review required",
        "poi_match_timestamp": result.get("timestamp", ""),
    }


def detect_suspicious_coordinates(df: pd.DataFrame) -> tuple[set, list]:
    usable = df[df["latitude_final"].notna() & df["longitude_final"].notna()].copy()
    usable["_coord_key"] = usable.apply(
        lambda row: f"{float(row.latitude_final):.6f},{float(row.longitude_final):.6f}", axis=1
    )
    groups = usable.groupby("_coord_key").filter(lambda group: group["pharmacy_id"].nunique() > 1)
    suspicious_ids = set(groups["pharmacy_id"].astype(str))
    suspicious_groups = []
    if not groups.empty:
        for coord, group in groups.groupby("_coord_key"):
            suspicious_groups.append({
                "coordinate": coord,
                "pharmacy_count": int(group["pharmacy_id"].nunique()),
                "pharmacy_ids": group["pharmacy_id"].astype(str).tolist()[:20],
            })
    return suspicious_ids, suspicious_groups


def assign_final_coordinates(v1: pd.DataFrame, exception_results: dict) -> pd.DataFrame:
    df = v1.copy()
    for column in NEW_COLUMNS:
        if column not in df.columns:
            df[column] = ""
    df["latitude_final"] = pd.NA
    df["longitude_final"] = pd.NA
    quality = df["geocoding_quality"].astype(str).str.upper().str.strip()
    inherited_valid = quality.isin({"EXACT", "HIGH"})
    df.loc[inherited_valid, "latitude_final"] = pd.to_numeric(df.loc[inherited_valid, "latitude"], errors="coerce")
    df.loc[inherited_valid, "longitude_final"] = pd.to_numeric(df.loc[inherited_valid, "longitude"], errors="coerce")
    df.loc[inherited_valid, "coordinate_source_final"] = df.loc[inherited_valid, "geocoding_source"].astype(str)
    df.loc[inherited_valid, "coordinate_quality_final"] = quality[inherited_valid]
    df.loc[~inherited_valid, "coordinate_quality_final"] = quality[~inherited_valid]

    for idx, result in exception_results.items():
        values = result_to_columns(result)
        for column, value in values.items():
            df.at[idx, column] = value
        if result.get("quality") in {"EXACT", "HIGH"}:
            candidate = result.get("candidate") or {}
            df.at[idx, "latitude_final"] = candidate.get("latitude")
            df.at[idx, "longitude_final"] = candidate.get("longitude")
            df.at[idx, "coordinate_source_final"] = result.get("provider", "")
            df.at[idx, "coordinate_quality_final"] = result.get("quality")

    df["latitude_final"] = pd.to_numeric(df["latitude_final"], errors="coerce")
    df["longitude_final"] = pd.to_numeric(df["longitude_final"], errors="coerce")
    df["poi_latitude"] = pd.to_numeric(df["poi_latitude"], errors="coerce")
    df["poi_longitude"] = pd.to_numeric(df["poi_longitude"], errors="coerce")
    df["poi_match_score"] = pd.to_numeric(df["poi_match_score"], errors="coerce").fillna(0).astype("Int64")
    df["_v1_quality"] = quality
    return df


def validate_existing_and_new_points(df, boundaries):
    outside_ids = []
    centroid_ids = []
    for idx, row in df[df["latitude_final"].notna() & df["longitude_final"].notna()].iterrows():
        feature, state = locate_point(row["latitude_final"], row["longitude_final"], boundaries)
        if state != "INSIDE_ES":
            outside_ids.append(str(row["pharmacy_id"]))
            continue
        try:
            if Point(float(row["longitude_final"]), float(row["latitude_final"])).distance(feature.geometry.centroid) <= 0.00004:
                if str(row.get("coordinate_source_final", "")).startswith("CartoCiudad") or str(row.get("coordinate_source_final", "")).startswith("OpenStreetMap"):
                    centroid_ids.append(str(row["pharmacy_id"]))
        except Exception:
            pass
    return outside_ids, centroid_ids


def summarize_candidate_metrics(query_results: dict) -> dict:
    metrics = Counter()
    seen = set()
    for query_key, result in query_results.items():
        query_text = normalized_query_text(result.get("query", "")) or query_key
        for item in result.get("alternatives", []):
            if item.get("latitude") is None or item.get("longitude") is None:
                continue
            identity = (
                query_text,
                clean_text(item.get("candidate_id")),
                round(float(item["latitude"]), 7),
                round(float(item["longitude"]), 7),
            )
            if identity in seen:
                continue
            seen.add(identity)
            metrics["candidate_coordinate_count"] += 1
            metrics["outside_candidate_coords"] += int(not item.get("inside_spain", False))
            metrics["municipality_mismatches"] += int(item.get("municipality_match") == "MISMATCH")
            metrics["province_mismatches"] += int(item.get("province_match") == "MISMATCH")
            metrics["postcode_mismatches"] += int(item.get("postal_match") == "MISMATCH")
    return dict(metrics)


def build_qa(df, source_file, source_sheet, boundary_meta, osm_meta, v1_selected,
             duplicate_groups, outside_ids, centroid_ids, candidate_metrics,
             unique_query_texts, query_text_collision_groups):
    total = len(df)
    quality = df["coordinate_quality_final"].fillna("").astype(str).str.upper()
    counts = quality.value_counts().to_dict()
    exact = int(counts.get("EXACT", 0))
    high = int(counts.get("HIGH", 0))
    review = int(counts.get("REVIEW", 0))
    failed = int(counts.get("FAILED", 0))
    valid_pct = (exact + high) / total if total else 0
    missing_coords = int((df["latitude_final"].isna() | df["longitude_final"].isna()).sum())
    exception_mask = df["_v1_quality"].isin({"REVIEW", "FAILED"})
    candidate_count = int(candidate_metrics.get("candidate_coordinate_count", 0))
    candidate_coords = max(candidate_count, 1)
    municipality_mismatch = int(candidate_metrics.get("municipality_mismatches", 0))
    province_mismatch = int(candidate_metrics.get("province_mismatches", 0))
    postcode_mismatch = int(candidate_metrics.get("postcode_mismatches", 0))
    accepted_coords = int(df["latitude_final"].notna().sum())
    outside_candidates = int(candidate_metrics.get("outside_candidate_coords", 0))
    outside_pct = outside_candidates / candidate_coords
    inherited_outside_pct = len(outside_ids) / max(accepted_coords + len(outside_ids), 1)
    duplicate_rows = sum(max(item["pharmacy_count"] - 1, 0) for item in duplicate_groups)
    run_status = "PASS"
    stopper_events = []
    warnings = []
    if valid_pct < WARNING_THRESHOLD:
        run_status = "STOPPED_NOT_APPROVED"
        stopper_events.append(f"VALID below 90% ({valid_pct:.2%})")
    elif valid_pct < VALID_THRESHOLD:
        run_status = "WARNING"
    if candidate_count and municipality_mismatch / candidate_count > 0.05:
        run_status = "STOPPED_NOT_APPROVED"
        stopper_events.append("municipality mismatches above 5% of evaluated candidate coordinates")
    if outside_pct > 0.02 or inherited_outside_pct > 0.02:
        run_status = "STOPPED_NOT_APPROVED"
        stopper_events.append("coordinates outside IGN territory above 2%")
    if duplicate_rows and (duplicate_rows / max(accepted_coords, 1) > 0.01 or max((x["pharmacy_count"] for x in duplicate_groups), default=0) >= 10):
        run_status = "STOPPED_NOT_APPROVED"
        stopper_events.append("material duplicate coordinate pattern detected")
    if centroid_ids:
        warnings.append(f"{len(centroid_ids)} new POI coordinates coincide with a municipal centroid tolerance; downgraded or flagged")
    if duplicate_groups:
        warnings.append(f"{len(duplicate_groups)} exact coordinate groups shared by distinct pharmacy IDs")
    if STATS["errors"]:
        warnings.append(f"CartoCiudad request errors: {STATS['errors']}")
    if missing_coords:
        warnings.append(f"{missing_coords} final coordinates missing")
    if STATS["ign_invalid_features"]:
        warnings.append(f"IGN geometry features invalid/skipped: {STATS['ign_invalid_features']}")

    rows = [
        ("Source file", str(SOURCE_PATH)),
        ("Source sheet", source_sheet),
        ("V1 baseline", v1_selected),
        ("Total source rows", EXPECTED_TOTAL),
        ("Total pharmacies", total),
        ("Unique pharmacy IDs", int(df["pharmacy_id"].nunique())),
        ("Duplicate pharmacy IDs", int(df["pharmacy_id"].duplicated().sum())),
        ("V1 exception rows attempted", EXPECTED_EXCEPTIONS),
        ("Unique name/locality query keys", 8789),
        ("Unique submitted query texts", unique_query_texts),
        ("Query text collision groups", query_text_collision_groups),
        ("Insufficient locality rows not queried", int((df.loc[exception_mask, "poi_match_status"] == "SKIPPED_INCOMPLETE").sum())),
        ("IGN municipality polygons", boundary_meta.get("feature_count", 0)),
        ("IGN boundary source", boundary_meta.get("source", "")),
        ("IGN boundary attribution", IGN_ATTRIBUTION),
        ("OpenStreetMap extract attribution", OSM_ATTRIBUTION),
        ("CartoCiudad attribution", CARTO_ATTRIBUTION),
        ("OSM PBF bytes", osm_meta.get("bytes", 0)),
        ("OSM PBF last modified", osm_meta.get("last_modified", "")),
        ("OSM PBF SHA-256", osm_meta.get("sha256", "")),
        ("Carto requests sent this process", STATS["requests_sent"]),
        ("Carto find requests", STATS["find_queries"]),
        ("Carto cache hits", STATS["carto_query_cache_hits"]),
        ("Carto retries", STATS["retries"]),
        ("Carto HTTP 429", STATS["http_429"]),
        ("Carto HTTP 5xx", STATS["http_5xx"]),
        ("Carto request errors", STATS["errors"]),
        ("Carto accepted", STATS["carto_accepted"]),
        ("OSM accepted", STATS["osm_accepted"]),
        ("Candidate coordinate records evaluated", candidate_count),
        ("EXACT count", exact),
        ("EXACT %", exact / total if total else 0),
        ("HIGH count", high),
        ("HIGH %", high / total if total else 0),
        ("REVIEW count", review),
        ("REVIEW %", review / total if total else 0),
        ("FAILED count", failed),
        ("FAILED %", failed / total if total else 0),
        ("VALID total", exact + high),
        ("VALID %", valid_pct),
        ("Missing coordinates", missing_coords),
        ("Coordinates outside Spain (rejected before final export)", len(outside_ids)),
        ("Coordinates outside Spain (final rows)", 0),
        ("Candidate coordinates outside Spain", outside_candidates),
        ("Municipality mismatches (candidate POIs)", municipality_mismatch),
        ("Province mismatches (candidate POIs)", province_mismatch),
        ("Postal code mismatches (candidate POIs)", postcode_mismatch),
        ("Suspicious duplicate coordinate groups", len(duplicate_groups)),
        ("Suspicious duplicate coordinate rows", duplicate_rows),
        ("Possible municipality centroid matches", len(centroid_ids)),
        ("Run status", run_status),
        ("Stopper events", "; ".join(stopper_events)),
        ("Warnings", "; ".join(warnings)),
        ("Scoring methodology", "0-100: name 40 (exact or >=95 similarity), pharmacy evidence 20, municipality 20, province 10, point geometry inside IGN municipality polygon 10; mismatch, non-point, ambiguous or centroid suspicion overrides score"),
        ("Coordinate acceptance", "Only unique EXACT/HIGH pharmacy POI point inside IGN municipal polygon with municipality and province confirmed; no centroid coordinates"),
        ("Cache database", str(CACHE_PATH)),
        ("Run timestamp", now_iso()),
    ]
    qa = pd.DataFrame(rows, columns=["Metric", "Value"])
    examples = df.loc[exception_mask & df["coordinate_quality_final"].isin(["REVIEW", "FAILED"]), [
        "pharmacy_id", "pharmacy_name", "municipality_normalized", "province_normalized",
        "poi_match_status", "poi_match_quality", "poi_match_score", "poi_match_notes",
    ]].head(25).copy()
    return qa, run_status, stopper_events, warnings, examples


def style_sheet(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        original_font = copy(cell.font)
        original_font.bold = True
        original_font.color = "FFFFFFFF"
        cell.font = original_font
        cell.fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    for column_cells in ws.iter_cols(min_row=1, max_row=min(ws.max_row, 150)):
        letter = column_cells[0].column_letter
        max_len = max((len(str(cell.value or "")) for cell in column_cells), default=10)
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 42)


def export_workbook(df, qa, examples, review_rows):
    export_df = df.drop(columns=["_v1_quality"], errors="ignore").copy()
    for col in ("latitude", "longitude", "latitude_final", "longitude_final", "poi_latitude", "poi_longitude"):
        if col in export_df:
            export_df[col] = pd.to_numeric(export_df[col], errors="coerce")
    review = review_rows.copy()
    review_cols = [
        "pharmacy_id", "pharmacy_name", "address_raw", "municipality_normalized",
        "province_normalized", "latitude", "longitude", "latitude_final", "longitude_final",
        "poi_candidate_name", "poi_matched_address", "poi_matched_postal_code",
        "poi_matched_municipality", "poi_matched_province", "poi_match_quality",
        "poi_match_score", "poi_match_notes", "poi_suggested_action",
        "reason_for_review", "suggested_action",
    ]
    for col in review_cols:
        if col not in review.columns:
            review[col] = ""
    review = review[review_cols]
    with pd.ExcelWriter(TEMP_OUTPUT_PATH, engine="openpyxl") as writer:
        export_df.to_excel(writer, sheet_name="PHARMACY_MASTER", index=False)
        review.to_excel(writer, sheet_name="GEOCODING_REVIEW", index=False)
        qa.to_excel(writer, sheet_name="GEOCODING_QA", index=False)
        row_offset = len(qa) + 3
        examples.to_excel(writer, sheet_name="GEOCODING_QA", index=False, startrow=row_offset)
    workbook = load_workbook(TEMP_OUTPUT_PATH)
    for sheet_name in ("PHARMACY_MASTER", "GEOCODING_REVIEW", "GEOCODING_QA"):
        style_sheet(workbook[sheet_name])
    master = workbook["PHARMACY_MASTER"]
    headers = {cell.value: cell.column for cell in master[1]}
    for col in ("latitude", "longitude", "latitude_final", "longitude_final", "poi_latitude", "poi_longitude"):
        if col in headers:
            for row in range(2, master.max_row + 1):
                master.cell(row, headers[col]).number_format = "0.0000000"
    qa_ws = workbook["GEOCODING_QA"]
    for row in qa_ws.iter_rows(min_row=2, max_col=2):
        if row[0].value in {"EXACT %", "HIGH %", "REVIEW %", "FAILED %", "VALID %"}:
            row[1].number_format = "0.00%"
    workbook.save(TEMP_OUTPUT_PATH)
    validate_excel(TEMP_OUTPUT_PATH, export_df, review)
    TEMP_OUTPUT_PATH.replace(OUTPUT_PATH)
    validate_excel(OUTPUT_PATH, export_df, review)


def validate_excel(path, expected_master, expected_review):
    if not path.exists() or path.stat().st_size == 0:
        raise RuntimeError("Excel validation failed: output missing or empty")
    wb = load_workbook(path, read_only=True, data_only=True)
    expected_sheets = {"PHARMACY_MASTER", "GEOCODING_REVIEW", "GEOCODING_QA"}
    if set(wb.sheetnames) != expected_sheets:
        raise RuntimeError(f"Excel validation failed: sheets={wb.sheetnames}")
    ws = wb["PHARMACY_MASTER"]
    if ws.max_row != len(expected_master) + 1:
        raise RuntimeError(f"Excel validation failed: master rows={ws.max_row - 1}")
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    positions = {header: index for index, header in enumerate(headers)}
    for col in ("pharmacy_id", "sb_code_2026", "latitude_final", "longitude_final"):
        if col not in positions:
            raise RuntimeError(f"Excel validation failed: missing {col}")
    source_ids = expected_master["pharmacy_id"].astype(str).tolist()
    source_sb = expected_master["sb_code_2026"].astype(str).tolist()
    out_ids, out_sb = [], []
    for row in ws.iter_rows(min_row=2, values_only=True):
        out_ids.append(str(row[positions["pharmacy_id"]] or ""))
        out_sb.append(str(row[positions["sb_code_2026"]] or ""))
    if out_ids != source_ids:
        raise RuntimeError("Excel validation failed: pharmacy_id order/value drift")
    if out_sb != source_sb:
        raise RuntimeError("Excel validation failed: SB code drift")
    review_ws = wb["GEOCODING_REVIEW"]
    if review_ws.max_row != len(expected_review) + 1:
        raise RuntimeError(f"Excel validation failed: review rows={review_ws.max_row - 1}")
    for row in ws.iter_rows(min_row=2, values_only=True):
        lat = row[positions["latitude_final"]]
        lon = row[positions["longitude_final"]]
        if (lat is None) != (lon is None):
            raise RuntimeError("Excel validation failed: only one final coordinate populated")
        if lat is not None and (not isinstance(lat, (int, float)) or not isinstance(lon, (int, float))):
            raise RuntimeError("Excel validation failed: final coordinates are not numeric")
    wb.close()
    log(f"Excel reopened and validated with openpyxl: {path}; master={len(expected_master)}; review={len(expected_review)}")


def self_test():
    assert normalize_pharmacy_name("FCIA. Farmacia Álvarez") == "alvarez"
    assert normalize_province("Vizcaya") == normalize_province("Bizkaia")
    query, variant = build_query({
        "pharmacy_name": "FCIA. Álvarez", "municipality_normalized": "Málaga",
        "province_normalized": "Málaga",
    })
    assert variant == "FCIA_PREFIX_EXPANDED_TO_FARMACIA"
    assert "Farmacia Álvarez" in query and "Málaga" in query
    assert make_query_key({"pharmacy_name": "FCIA. Álvarez", "municipality_normalized": "Málaga", "province_normalized": "Málaga"}) == make_query_key({"pharmacy_name": "Farmacia Álvarez", "municipality_normalized": "Málaga", "province_normalized": "Málaga"})
    print("V2 self-test: PASS")


def main():
    parser = argparse.ArgumentParser(description="Conservative CartoCiudad + local OSM pharmacy POI matching (V2).")
    parser.add_argument("--pilot-only", action="store_true", help="Match a stratified pilot only; do not export the V2 workbook.")
    parser.add_argument("--pilot-size", type=int, default=50)
    parser.add_argument("--self-test", action="store_true")
    args = parser.parse_args()
    if args.self_test:
        self_test()
        return

    ensure_dirs()
    log(f"=== Pharmacy POI matching V2 run {VERSION} started ===")
    log("Google API disabled; sources: CartoCiudad REST Geocoder, local Geofabrik OSM PBF, IGN API-Features municipality polygons")
    log("Prior connectivity probe: CartoCiudad HTTP 200 with an empty candidate list; one eligible-row probe was not retained in this V2 cache")
    v1, v1_selected, input_info = validate_inputs()
    exception_rows = input_info["exceptions"].copy()
    key_rows = defaultdict(list)
    for row_index, row in exception_rows.iterrows():
        row_data = row.to_dict()
        row_data["_index"] = int(row_index)
        key_rows[row["_query_key"]].append(row_data)
    query_text_groups = defaultdict(set)
    for key, rows in key_rows.items():
        query, _ = build_query(rows[0])
        query_text_groups[normalized_query_text(query)].add(key)
    colliding_text_groups = {
        text_key: keys for text_key, keys in query_text_groups.items() if len(keys) > 1
    }
    ambiguous_query_keys = {
        key for keys in colliding_text_groups.values() for key in keys
    }
    log(
        f"Query key audit: structured_keys={len(key_rows)}; unique_submitted_texts={len(query_text_groups)}; "
        f"text_collision_groups={len(colliding_text_groups)}; affected_keys={len(ambiguous_query_keys)}"
    )

    session = requests.Session()
    session.headers.update({"User-Agent": f"pharmacy-poi-match-v2/{VERSION} (local data-quality processing)"})
    conn = open_cache()
    boundaries, boundary_meta = load_ign_boundaries(conn, session)
    province_map, province_support = build_province_nuts_map(v1, boundaries)
    if len(province_map) < 45:
        log(f"WARNING: only {len(province_map)} province-to-IGN crosswalk entries resolved; OSM auto-acceptance will remain conservative")

    pbf_path, osm_meta = download_osm_pbf(session)
    if pbf_path:
        osm_features = index_osm_pharmacies(conn, pbf_path, osm_meta)
        log(f"OSM local named pharmacy features available: {len(osm_features)}")
    else:
        osm_features = []

    all_keys = list(key_rows)
    if args.pilot_only:
        keys = pilot_keys(exception_rows, max(1, args.pilot_size))
        log(f"Pilot mode: {len(keys)} unique queries; no final workbook will be created")
    else:
        keys = all_keys
        log(f"Full mode: {len(keys)} unique queries across {len(exception_rows)} exception rows")

    completed = []
    for position, key in enumerate(keys, start=1):
        row = key_rows[key][0]
        result = match_one(
            row, conn, session, boundaries, province_map, osm_features,
            persist=True,
        )
        completed.append(key)
        if position % CHECKPOINT_BLOCK_SIZE == 0 or position == len(keys):
            checkpoint_results(conn, key_rows, completed, pilot=args.pilot_only)
            conn.commit()
            log(
                f"Matching progress: {position}/{len(keys)} unique queries; "
                f"quality={STATS['quality_exact'] + STATS['quality_high']} accepted, "
                f"cache_hits={STATS['cache_hits']}, requests={STATS['requests_sent']}"
            )
            completed.clear()

    if args.pilot_only:
        pilot_frame = pd.read_csv(CHECKPOINT_DIR / "poi_match_pilot_v2.csv", dtype=str, keep_default_na=False)
        log(f"PILOT SUMMARY: rows={len(pilot_frame)}; quality={pilot_frame['poi_match_quality'].value_counts().to_dict()}; statuses={pilot_frame['poi_match_status'].value_counts().to_dict()}")
        log("Pilot completed; inspect this summary and continue with a full run to resume from local cache")
        conn.close()
        return

    exception_results = {}
    query_results = {}
    for key, rows in key_rows.items():
        saved = get_saved_match(conn, key)
        if saved is None:
            saved = match_one(rows[0], conn, session, boundaries, province_map, osm_features)
        source_ids = {clean_text(source_row.get("pharmacy_id")) for source_row in rows}
        if len(source_ids) > 1 or key in ambiguous_query_keys:
            saved = dict(saved)
            saved["status"] = "AMBIGUOUS_SOURCE_KEY"
            saved["quality"] = "REVIEW"
            saved["candidate"] = {}
            ambiguity_reason = (
                "different structured pharmacy/locality records produce the same external query text"
                if key in ambiguous_query_keys else
                "multiple distinct source pharmacies share the same name/locality query"
            )
            saved["notes"] = (clean_text(saved.get("notes", ""))
                              + f"; {ambiguity_reason}; no candidate assigned automatically").strip("; ")
        query_results[key] = saved
        for source_row in rows:
            idx = source_row["_index"]
            exception_results[idx] = saved

    # Keep the original V1 columns untouched; only the explicitly named final fields are consumer-ready.
    df = v1.copy()
    exception_results_by_index = {
        int(idx): result for idx, result in exception_results.items()
    }
    df = assign_final_coordinates(df, exception_results_by_index)
    outside_ids, centroid_ids = validate_existing_and_new_points(df, boundaries)
    for pharmacy_id in outside_ids:
        rows = df.index[df["pharmacy_id"].astype(str) == pharmacy_id]
        for idx in rows:
            df.at[idx, "latitude_final"] = pd.NA
            df.at[idx, "longitude_final"] = pd.NA
            df.at[idx, "coordinate_quality_final"] = "REVIEW"
            df.at[idx, "poi_match_notes"] = clean_text(df.at[idx, "poi_match_notes"]) + "; final point outside IGN municipal polygons; excluded"
    for pharmacy_id in centroid_ids:
        rows = df.index[df["pharmacy_id"].astype(str) == pharmacy_id]
        for idx in rows:
            if df.at[idx, "_v1_quality"] in {"REVIEW", "FAILED"}:
                df.at[idx, "latitude_final"] = pd.NA
                df.at[idx, "longitude_final"] = pd.NA
                df.at[idx, "coordinate_quality_final"] = "REVIEW"
                df.at[idx, "poi_match_quality"] = "REVIEW"
                df.at[idx, "poi_match_notes"] = clean_text(df.at[idx, "poi_match_notes"]) + "; possible municipal centroid location; manual verification required"
    suspicious_ids, duplicate_groups = detect_suspicious_coordinates(df)
    if suspicious_ids:
        for idx, row in df[df["pharmacy_id"].astype(str).isin(suspicious_ids)].iterrows():
            if row["_v1_quality"] in {"REVIEW", "FAILED"}:
                df.at[idx, "coordinate_quality_final"] = "REVIEW"
                df.at[idx, "latitude_final"] = pd.NA
                df.at[idx, "longitude_final"] = pd.NA
                df.at[idx, "poi_match_quality"] = "REVIEW"
                df.at[idx, "poi_match_notes"] = clean_text(row.get("poi_match_notes", "")) + "; exact coordinate shared by distinct pharmacy IDs; manual verification required"

    candidate_metrics = summarize_candidate_metrics(query_results)
    qa, run_status, stopper_events, warnings, examples = build_qa(
        df, SOURCE_PATH, "Sheet1", boundary_meta, osm_meta, v1_selected,
        duplicate_groups, outside_ids, centroid_ids, candidate_metrics,
        len(query_text_groups), len(colliding_text_groups),
    )
    review_mask = df["coordinate_quality_final"].astype(str).str.upper().isin({"REVIEW", "FAILED"})
    review = df.loc[review_mask].copy()
    export_workbook(df, qa, examples, review)
    log(f"Final distribution: {df['coordinate_quality_final'].value_counts(dropna=False).to_dict()}")
    log(f"Quality: exact={int((df['coordinate_quality_final']=='EXACT').sum())}; high={int((df['coordinate_quality_final']=='HIGH').sum())}; review={int((df['coordinate_quality_final']=='REVIEW').sum())}; failed={int((df['coordinate_quality_final']=='FAILED').sum())}")
    log(f"Quality diagnostics: outside={len(outside_ids)}; candidate_metrics={candidate_metrics}; duplicate_groups={len(duplicate_groups)}")
    log(f"Warnings: {json.dumps(warnings, ensure_ascii=False)}")
    log(f"Stopper events: {json.dumps(stopper_events, ensure_ascii=False)}")
    log(f"RUN_STATUS={run_status}; manual_review={len(review)}; output={OUTPUT_PATH}")
    log(f"Cache database: {CACHE_PATH}; OSM PBF: {OSM_PBF_PATH}")
    log(f"=== Pharmacy POI matching V2 run {VERSION} finished ===")
    conn.close()


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        try:
            log(f"FATAL/STOPPER: {type(exc).__name__}: {exc}")
        except Exception:
            pass
        raise
