from __future__ import annotations

import math
import re
from dataclasses import dataclass
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


BASE_DIR = Path(__file__).resolve().parents[2]
PHARMACY_DIR = BASE_DIR / "pharmacy"
RX_2025_PATH = PHARMACY_DIR / "rx" / "RX_2025.xlsx"
RX_2026_PATH = PHARMACY_DIR / "rx" / "RX 2026.xlsx"
STRUCTURE_ANALYSIS_PATH = BASE_DIR / "pharmacy/input/iqvia/Análisis de la estructura SB2025 - SB2026_No Onekey.xlsx"
STRUCTURE_2026_PATH = BASE_DIR / "pharmacy/input/iqvia/Estructura SB_No Onekey_2026.xlsx"
OUTPUT_PATH = PHARMACY_DIR / "output" / "RX_2024_2026_SB2026.xlsx"

KEY_COLS = ["SB CODE", "Manufacturer", "Measure"]
EXPECTED_MANUFACTURERS = ["ESSITY", "RESTO"]
EXPECTED_MEASURES = ["Euros", "Units"]
OLD_CODES = {
    "020210",
    "020280",
    "081134",
    "260190",
    "260250",
    "260300",
    "260290",
    "330660",
    "331050",
}
NEW_CODES = {
    "020600",
    "081461",
    "081464",
    "260340",
    "260330",
    "260350",
    "260360",
    "331160",
    "331170",
}

UNION_MAPPINGS = {"020210": "020600", "020280": "020600"}
SPLIT_MAPPINGS = {"081134": {"081461": 0.5, "081464": 0.5}}
ONE_TO_ONE_MAPPINGS = {
    "260190": "260340",
    "260250": "260330",
    "260300": "260350",
    "260290": "260360",
    "330660": "331160",
    "331050": "331170",
}


@dataclass
class LoadedSales:
    path: Path
    sheet_name: str
    raw: pd.DataFrame
    usable: pd.DataFrame
    month_cols: list[str]
    month_map: dict[str, pd.Timestamp]


def normalize_sb_code(value) -> str:
    if pd.isna(value):
        return ""
    text = str(value).strip()
    if re.fullmatch(r"\d+\.0", text):
        text = text[:-2]
    digits = re.sub(r"\D", "", text)
    if digits and len(digits) <= 6:
        return digits.zfill(6)
    return text


def normalize_manufacturer(value) -> str:
    return "" if pd.isna(value) else str(value).strip().upper()


def normalize_measure(value) -> str:
    if pd.isna(value):
        return ""
    text = str(value).strip()
    low = text.lower()
    if "uro" in low or "€" in text:
        return "Euros"
    if "unit" in low:
        return "Units"
    return text


def parse_month_label(label) -> pd.Timestamp | None:
    if isinstance(label, pd.Timestamp):
        return pd.Timestamp(label.year, label.month, 1)
    if hasattr(label, "year") and hasattr(label, "month"):
        return pd.Timestamp(label.year, label.month, 1)
    text = str(label).strip()
    if not text:
        return None
    month_names = {
        "jan": 1,
        "feb": 2,
        "mar": 3,
        "apr": 4,
        "may": 5,
        "jun": 6,
        "jul": 7,
        "aug": 8,
        "sep": 9,
        "oct": 10,
        "nov": 11,
        "dec": 12,
        "ene": 1,
        "abr": 4,
        "ago": 8,
        "dic": 12,
    }
    match = re.fullmatch(r"([A-Za-zÁÉÍÓÚáéíóú]{3})[\s\-_/]*(\d{2,4})", text)
    if match:
        mon = month_names.get(match.group(1).lower()[:3])
        year = int(match.group(2))
        year = 2000 + year if year < 100 else year
        if mon:
            return pd.Timestamp(year, mon, 1)
    match = re.fullmatch(r"(20\d{2})(0[1-9]|1[0-2])", text)
    if match:
        return pd.Timestamp(int(match.group(1)), int(match.group(2)), 1)
    match = re.fullmatch(r"(0?[1-9]|1[0-2])[/\-](20\d{2}|\d{2})", text)
    if match:
        year = int(match.group(2))
        year = 2000 + year if year < 100 else year
        return pd.Timestamp(year, int(match.group(1)), 1)
    parsed = pd.to_datetime(text, errors="coerce")
    if pd.notna(parsed) and 2000 <= parsed.year <= 2035:
        return pd.Timestamp(parsed.year, parsed.month, 1)
    return None


def month_label(month: pd.Timestamp) -> str:
    return month.strftime("%b-%y")


def expected_months() -> list[pd.Timestamp]:
    return list(pd.date_range("2024-01-01", "2026-06-01", freq="MS"))


def detect_month_columns(df: pd.DataFrame) -> tuple[list[str], dict[str, pd.Timestamp]]:
    month_map: dict[str, pd.Timestamp] = {}
    for col in df.columns:
        parsed = parse_month_label(col)
        if parsed is not None:
            month_map[col] = parsed
    month_cols = sorted(month_map, key=month_map.get)
    return month_cols, month_map


def inspect_excel(path: Path) -> dict:
    xls = pd.ExcelFile(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for sheet in xls.sheet_names:
        ws = workbook[sheet]
        sample = pd.read_excel(path, sheet_name=sheet, nrows=5, dtype=object)
        full = pd.read_excel(path, sheet_name=sheet, dtype=object)
        month_cols, month_map = detect_month_columns(full)
        sheets.append(
            {
                "sheet": sheet,
                "openpyxl_rows": ws.max_row,
                "openpyxl_cols": ws.max_column,
                "pandas_rows": len(full),
                "pandas_cols": len(full.columns),
                "headers": [str(c) for c in full.columns],
                "first_rows": sample.to_dict("records"),
                "dtypes": {str(k): str(v) for k, v in full.dtypes.items()},
                "month_cols": [str(c) for c in month_cols],
                "month_min": month_label(min(month_map.values())) if month_map else "",
                "month_max": month_label(max(month_map.values())) if month_map else "",
                "sb_distinct": (
                    full["SB CODE"].map(normalize_sb_code).nunique()
                    if "SB CODE" in full.columns
                    else None
                ),
                "manufacturers": (
                    sorted(full["MANUFACTURER"].dropna().map(normalize_manufacturer).unique().tolist())
                    if "MANUFACTURER" in full.columns
                    else []
                ),
                "measures": (
                    sorted(full["Measure"].dropna().map(normalize_measure).unique().tolist())
                    if "Measure" in full.columns
                    else []
                ),
            }
        )
    return {"path": str(path), "sheets": sheets}


def print_inspection(inspection: dict) -> None:
    print("\n" + "=" * 100)
    print(f"FILE: {inspection['path']}")
    for sheet in inspection["sheets"]:
        print(f"-- SHEET: {sheet['sheet']}")
        print(f"   dimensions: {sheet['openpyxl_rows']} rows x {sheet['openpyxl_cols']} cols")
        print(f"   pandas shape: {sheet['pandas_rows']} rows x {sheet['pandas_cols']} cols")
        print(f"   headers: {sheet['headers'][:35]}")
        print(f"   month columns: {len(sheet['month_cols'])} ({sheet['month_min']} -> {sheet['month_max']})")
        print(f"   SB CODE distinct: {sheet['sb_distinct']}")
        print(f"   Manufacturers: {sheet['manufacturers']}")
        print(f"   Measures: {sheet['measures']}")
        print(f"   first rows: {sheet['first_rows'][:2]}")


def load_sales_file(path: Path) -> LoadedSales:
    xls = pd.ExcelFile(path)
    sheet_name = xls.sheet_names[0]
    raw = pd.read_excel(path, sheet_name=sheet_name, dtype=object)
    month_cols, month_map = detect_month_columns(raw)
    required = {"SB CODE", "MANUFACTURER", "Measure"}
    missing = required - set(raw.columns)
    if missing:
        raise AssertionError(f"{path.name} missing required columns: {missing}")
    usable = raw.copy()
    usable["SB CODE"] = usable["SB CODE"].map(normalize_sb_code)
    usable["Manufacturer"] = usable["MANUFACTURER"].map(normalize_manufacturer)
    usable["Measure"] = usable["Measure"].map(normalize_measure)
    for col in month_cols:
        usable[col] = pd.to_numeric(usable[col], errors="coerce")
    return LoadedSales(path, sheet_name, raw, usable, month_cols, month_map)


def to_long(data: LoadedSales, months: list[pd.Timestamp], keep_expected_manufacturers: bool = True) -> pd.DataFrame:
    selected_cols = [col for col in data.month_cols if data.month_map[col] in months]
    df = data.usable[["SB CODE", "Manufacturer", "Measure"] + selected_cols].copy()
    if keep_expected_manufacturers:
        df = df[df["Manufacturer"].isin(EXPECTED_MANUFACTURERS)].copy()
    long = df.melt(id_vars=KEY_COLS, value_vars=selected_cols, var_name="month_source_col", value_name="value")
    long["Month"] = long["month_source_col"].map(data.month_map)
    long["value"] = pd.to_numeric(long["value"], errors="coerce")
    return long[KEY_COLS + ["Month", "value"]]


def aggregate_long(long: pd.DataFrame) -> pd.DataFrame:
    return long.groupby(KEY_COLS + ["Month"], as_index=False, dropna=False)["value"].sum(min_count=1)


def transform_2025_to_2026(long: pd.DataFrame) -> pd.DataFrame:
    pieces = []
    unchanged = long[~long["SB CODE"].isin(OLD_CODES)].copy()
    pieces.append(unchanged)

    union_part = long[long["SB CODE"].isin(UNION_MAPPINGS)].copy()
    union_part["SB CODE"] = union_part["SB CODE"].map(UNION_MAPPINGS)
    pieces.append(union_part)

    for old, targets in SPLIT_MAPPINGS.items():
        source = long[long["SB CODE"] == old]
        for new, factor in targets.items():
            split = source.copy()
            split["SB CODE"] = new
            split["value"] = split["value"] * factor
            pieces.append(split)

    one_to_one = long[long["SB CODE"].isin(ONE_TO_ONE_MAPPINGS)].copy()
    one_to_one["SB CODE"] = one_to_one["SB CODE"].map(ONE_TO_ONE_MAPPINGS)
    pieces.append(one_to_one)

    transformed = pd.concat(pieces, ignore_index=True)
    return aggregate_long(transformed)


def totals_by_month_manufacturer_measure(long: pd.DataFrame) -> pd.DataFrame:
    return long.groupby(["Month", "Manufacturer", "Measure"], as_index=False)["value"].sum(min_count=1)


def validate_market_conservation(before: pd.DataFrame, after: pd.DataFrame, label: str) -> pd.DataFrame:
    before_totals = totals_by_month_manufacturer_measure(before).rename(columns={"value": "before"})
    after_totals = totals_by_month_manufacturer_measure(after).rename(columns={"value": "after"})
    check = before_totals.merge(after_totals, on=["Month", "Manufacturer", "Measure"], how="outer")
    check["before"] = check["before"].fillna(0)
    check["after"] = check["after"].fillna(0)
    check["diff"] = check["after"] - check["before"]
    max_abs_diff = check["diff"].abs().max()
    assert max_abs_diff <= 1e-6, f"{label} market conservation failed. Max abs diff: {max_abs_diff}"
    return check


def assert_unique_key(df: pd.DataFrame, keys: list[str], label: str) -> None:
    duplicates = df.duplicated(keys).sum()
    assert duplicates == 0, f"{label} has {duplicates} duplicate rows for {keys}"


def validate_transformation_checks(before: pd.DataFrame, after: pd.DataFrame) -> dict[str, pd.DataFrame]:
    checks: dict[str, pd.DataFrame] = {}

    def values_for(frame: pd.DataFrame, code: str, value_name: str) -> pd.DataFrame:
        return frame[frame["SB CODE"] == code][["Month", "Manufacturer", "Measure", "value"]].rename(
            columns={"value": value_name}
        )

    alb = values_for(before, "020210", "old_020210").merge(
        values_for(before, "020280", "old_020280"), on=["Month", "Manufacturer", "Measure"], how="outer"
    )
    alb = alb.merge(values_for(after, "020600", "new_020600"), on=["Month", "Manufacturer", "Measure"], how="outer")
    alb[["old_020210", "old_020280", "new_020600"]] = alb[["old_020210", "old_020280", "new_020600"]].fillna(0)
    alb["expected"] = alb["old_020210"] + alb["old_020280"]
    alb["diff"] = alb["new_020600"] - alb["expected"]
    assert alb["diff"].abs().max() <= 1e-6, "Albacete union check failed"
    checks["Albacete"] = alb

    bcn = values_for(before, "081134", "old_081134")
    bcn_new = values_for(after, "081461", "new_081461").merge(
        values_for(after, "081464", "new_081464"), on=["Month", "Manufacturer", "Measure"], how="outer"
    )
    bcn = bcn.merge(bcn_new, on=["Month", "Manufacturer", "Measure"], how="outer")
    bcn[["old_081134", "new_081461", "new_081464"]] = bcn[["old_081134", "new_081461", "new_081464"]].fillna(0)
    bcn["new_sum"] = bcn["new_081461"] + bcn["new_081464"]
    bcn["diff"] = bcn["new_sum"] - bcn["old_081134"]
    assert bcn["diff"].abs().max() <= 1e-6, "Barcelona split check failed"
    checks["Barcelona"] = bcn

    one_rows = []
    for old, new in ONE_TO_ONE_MAPPINGS.items():
        merged = values_for(before, old, "old_value").merge(
            values_for(after, new, "new_value"), on=["Month", "Manufacturer", "Measure"], how="outer"
        )
        merged[["old_value", "new_value"]] = merged[["old_value", "new_value"]].fillna(0)
        merged["Old SB"] = old
        merged["New SB"] = new
        merged["diff"] = merged["new_value"] - merged["old_value"]
        one_rows.append(merged)
    one = pd.concat(one_rows, ignore_index=True)
    assert one["diff"].abs().max() <= 1e-6, "1:1 code change check failed"
    checks["One-to-one"] = one
    return checks


def compare_overlap(rx2025: LoadedSales, rx2026: LoadedSales) -> pd.DataFrame:
    months_2025 = set(rx2025.month_map.values())
    months_2026 = set(rx2026.month_map.values())
    overlap = sorted(months_2025 & months_2026)
    old_long = to_long(rx2025, overlap, keep_expected_manufacturers=True)
    new_long = to_long(rx2026, overlap, keep_expected_manufacturers=True)
    old_totals = totals_by_month_manufacturer_measure(old_long).rename(columns={"value": "RX_2025"})
    new_totals = totals_by_month_manufacturer_measure(new_long).rename(columns={"value": "RX_2026"})
    comp = old_totals.merge(new_totals, on=["Month", "Manufacturer", "Measure"], how="outer")
    comp[["RX_2025", "RX_2026"]] = comp[["RX_2025", "RX_2026"]].fillna(0)
    comp["Diff_RX2026_minus_RX2025"] = comp["RX_2026"] - comp["RX_2025"]
    comp["Abs_Diff"] = comp["Diff_RX2026_minus_RX2025"].abs()
    return comp.sort_values(["Month", "Manufacturer", "Measure"])


def build_monthly_market_table(final_long: pd.DataFrame) -> pd.DataFrame:
    pivot = (
        final_long.pivot_table(
            index="Month",
            columns=["Manufacturer", "Measure"],
            values="value",
            aggfunc="sum",
            dropna=False,
        )
        .sort_index()
        .fillna(0)
    )
    result = pd.DataFrame({"Month": [month_label(m) for m in pivot.index]})
    for manufacturer in EXPECTED_MANUFACTURERS:
        for measure in EXPECTED_MEASURES:
            result[f"{manufacturer} {measure}"] = [
                pivot.loc[m, (manufacturer, measure)] if (manufacturer, measure) in pivot.columns else 0 for m in pivot.index
            ]
    result["Total Market Units"] = result["ESSITY Units"] + result["RESTO Units"]
    result["Total Market Euros"] = result["ESSITY Euros"] + result["RESTO Euros"]
    return result[
        [
            "Month",
            "ESSITY Units",
            "RESTO Units",
            "Total Market Units",
            "ESSITY Euros",
            "RESTO Euros",
            "Total Market Euros",
        ]
    ]


def build_transformation_audit() -> pd.DataFrame:
    rows = [
        ["020210", "020600", "UNION", 1, "Jan-24-Jun-24", "summed with 020280"],
        ["020280", "020600", "UNION", 1, "Jan-24-Jun-24", "summed with 020210"],
        ["081134", "081461", "SPLIT", 0.5, "Jan-24-Jun-24", "equal allocation"],
        ["081134", "081464", "SPLIT", 0.5, "Jan-24-Jun-24", "equal allocation"],
    ]
    for old, new in ONE_TO_ONE_MAPPINGS.items():
        region = "LA RIOJA" if old.startswith("26") else "ASTURIAS"
        rows.append([old, new, "1:1 CODE CHANGE", 1, "Jan-24-Jun-24", f"{region}; value unchanged"])
    return pd.DataFrame(
        rows,
        columns=["Old SB", "New SB", "Transformation Type", "Allocation Factor", "Period Affected", "Notes"],
    )


def build_validation_summary(
    rx2025: LoadedSales,
    rx2026: LoadedSales,
    final_data: pd.DataFrame,
    final_long: pd.DataFrame,
    checks: dict[str, pd.DataFrame],
    overlap: pd.DataFrame,
    market_check: pd.DataFrame,
    warnings: list[str],
) -> pd.DataFrame:
    old_found = sorted(set(final_data["SB CODE"]) & OLD_CODES)
    new_found = sorted(set(final_data["SB CODE"]) & NEW_CODES)
    dup_count = final_data.duplicated(KEY_COLS).sum()
    overlap_material = overlap[overlap["Abs_Diff"] > 1e-6]
    max_overlap = overlap["Abs_Diff"].max() if len(overlap) else 0
    summary_rows = [
        ("RX 2025 file used", str(rx2025.path)),
        ("RX 2026 file used", str(rx2026.path)),
        ("Structure analysis file inspected", str(STRUCTURE_ANALYSIS_PATH)),
        ("Structure 2026 file inspected", str(STRUCTURE_2026_PATH)),
        ("RX 2025 sheet used", rx2025.sheet_name),
        ("RX 2026 sheet used", rx2026.sheet_name),
        ("RX 2025 original rows", len(rx2025.raw)),
        ("RX 2026 original rows", len(rx2026.raw)),
        ("RX 2025 original SB distinct", rx2025.usable["SB CODE"].nunique()),
        ("RX 2026 original SB distinct", rx2026.usable["SB CODE"].nunique()),
        ("Range of months", "Jan-24 to Jun-26"),
        ("Final month count", len(expected_months())),
        ("Final rows", len(final_data)),
        ("Final Sanibricks", final_data["SB CODE"].nunique()),
        ("Manufacturers found in RX 2025", ", ".join(sorted(rx2025.usable["Manufacturer"].unique()))),
        ("Manufacturers found in RX 2026", ", ".join(sorted(rx2026.usable["Manufacturer"].unique()))),
        ("Manufacturers in output", ", ".join(sorted(final_data["Manufacturer"].unique()))),
        ("Measures found in RX 2025", ", ".join(sorted(rx2025.usable["Measure"].unique()))),
        ("Measures found in RX 2026", ", ".join(sorted(rx2026.usable["Measure"].unique()))),
        ("Measures in output", ", ".join(sorted(final_data["Measure"].unique()))),
        ("Duplicate SB CODE x Manufacturer x Measure rows", int(dup_count)),
        ("Old codes found in final", ", ".join(old_found) if old_found else "None"),
        ("New codes found in final", ", ".join(new_found)),
        ("Market conserved", "YES" if market_check["diff"].abs().max() <= 1e-6 else "NO"),
        ("Albacete check", "PASS" if checks["Albacete"]["diff"].abs().max() <= 1e-6 else "FAIL"),
        ("Barcelona check", "PASS" if checks["Barcelona"]["diff"].abs().max() <= 1e-6 else "FAIL"),
        (
            "La Rioja check",
            "PASS"
            if checks["One-to-one"].query("`Old SB`.str.startswith('26')", engine="python")["diff"].abs().max() <= 1e-6
            else "FAIL",
        ),
        (
            "Asturias check",
            "PASS"
            if checks["One-to-one"].query("`Old SB`.str.startswith('33')", engine="python")["diff"].abs().max() <= 1e-6
            else "FAIL",
        ),
        ("Overlap comparable months", f"{month_label(overlap['Month'].min())} to {month_label(overlap['Month'].max())}"),
        ("Overlap rows with differences", len(overlap_material)),
        ("Overlap max absolute difference", max_overlap),
        ("Warnings", " | ".join(warnings) if warnings else "None"),
    ]
    return pd.DataFrame(summary_rows, columns=["Metric", "Value"])


def export_excel(
    output_path: Path,
    final_data: pd.DataFrame,
    validation_summary: pd.DataFrame,
    monthly_market: pd.DataFrame,
    audit: pd.DataFrame,
    checks: dict[str, pd.DataFrame],
    overlap: pd.DataFrame,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        final_data.to_excel(writer, index=False, sheet_name="RX_DATA")
        validation_summary.to_excel(writer, index=False, sheet_name="VALIDATION", startrow=0)
        start = len(validation_summary) + 3
        monthly_market.to_excel(writer, index=False, sheet_name="VALIDATION", startrow=start)
        start += len(monthly_market) + 3
        overlap_export = overlap.copy()
        overlap_export["Month"] = overlap_export["Month"].map(month_label)
        overlap_export.to_excel(writer, index=False, sheet_name="VALIDATION", startrow=start)
        start += len(overlap_export) + 3
        for check_name, check_df in checks.items():
            export_df = check_df.copy()
            export_df["Month"] = export_df["Month"].map(month_label)
            export_df.insert(0, "Check", check_name)
            export_df.to_excel(writer, index=False, sheet_name="VALIDATION", startrow=start)
            start += len(export_df) + 3
        audit.to_excel(writer, index=False, sheet_name="SB_TRANSFORMATION_AUDIT")

    wb = load_workbook(output_path)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        for col_idx, column_cells in enumerate(ws.columns, start=1):
            max_len = 0
            for cell in column_cells[:250]:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 28)
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    cell.number_format = "#,##0.00" if not math.isclose(cell.value, round(cell.value), abs_tol=1e-9) else "#,##0"
    wb.save(output_path)


def validate_output_workbook(output_path: Path) -> dict:
    assert output_path.exists(), f"Missing output workbook: {output_path}"
    wb = load_workbook(output_path, read_only=True, data_only=True)
    expected_sheets = {"RX_DATA", "VALIDATION", "SB_TRANSFORMATION_AUDIT"}
    assert expected_sheets.issubset(set(wb.sheetnames)), f"Missing sheets: {expected_sheets - set(wb.sheetnames)}"
    df = pd.read_excel(output_path, sheet_name="RX_DATA", dtype={"SB CODE": str})
    month_cols, month_map = detect_month_columns(df)
    assert len(month_cols) == 30, f"RX_DATA has {len(month_cols)} month columns, expected 30"
    assert sorted(month_map.values()) == expected_months(), "RX_DATA months are not Jan-2024 through Jun-2026"
    assert not (set(df["SB CODE"].map(normalize_sb_code)) & OLD_CODES), "Old codes found in final workbook"
    assert df.duplicated(KEY_COLS).sum() == 0, "Final workbook key is not unique"
    return {
        "sheets": wb.sheetnames,
        "rows": len(df),
        "sanibricks": df["SB CODE"].map(normalize_sb_code).nunique(),
        "months": len(month_cols),
        "old_codes": sorted(set(df["SB CODE"].map(normalize_sb_code)) & OLD_CODES),
    }


def build() -> dict:
    print("Inspecting source Excel files...")
    for path in [RX_2025_PATH, RX_2026_PATH, STRUCTURE_ANALYSIS_PATH, STRUCTURE_2026_PATH]:
        if path.exists():
            print_inspection(inspect_excel(path))
        else:
            print(f"WARNING: Expected input not found at {path}")

    rx2025 = load_sales_file(RX_2025_PATH)
    rx2026 = load_sales_file(RX_2026_PATH)
    warnings = []
    extra_2025_manufacturers = sorted(set(rx2025.usable["Manufacturer"]) - set(EXPECTED_MANUFACTURERS))
    if extra_2025_manufacturers:
        rows = len(rx2025.usable[rx2025.usable["Manufacturer"].isin(extra_2025_manufacturers)])
        warnings.append(
            f"RX_2025 contains non-output Manufacturer values {extra_2025_manufacturers} in {rows} rows; "
            "these are province aggregate rows and are excluded from RX_DATA."
        )
    extra_2026_manufacturers = sorted(set(rx2026.usable["Manufacturer"]) - set(EXPECTED_MANUFACTURERS))
    if extra_2026_manufacturers:
        warnings.append(f"RX_2026 contains non-expected Manufacturer values {extra_2026_manufacturers}.")

    jan_jun_2024 = list(pd.date_range("2024-01-01", "2024-06-01", freq="MS"))
    jul_2024_jun_2026 = list(pd.date_range("2024-07-01", "2026-06-01", freq="MS"))
    all_months = expected_months()

    selected_2025_before = aggregate_long(to_long(rx2025, jan_jun_2024, keep_expected_manufacturers=True))
    selected_2025_after = transform_2025_to_2026(selected_2025_before)
    market_check = validate_market_conservation(selected_2025_before, selected_2025_after, "Jan-Jun 2024 transformation")
    checks = validate_transformation_checks(selected_2025_before, selected_2025_after)

    selected_2026 = aggregate_long(to_long(rx2026, jul_2024_jun_2026, keep_expected_manufacturers=True))
    final_long = aggregate_long(pd.concat([selected_2025_after, selected_2026], ignore_index=True))

    assert sorted(final_long["Month"].unique()) == all_months, "Final month coverage is incomplete or duplicated"
    assert set(final_long["Manufacturer"].unique()) <= set(EXPECTED_MANUFACTURERS), "Unexpected Manufacturer in final data"
    assert set(final_long["Measure"].unique()) <= set(EXPECTED_MEASURES), "Unexpected Measure in final data"

    final_data = final_long.pivot_table(
        index=KEY_COLS,
        columns="Month",
        values="value",
        aggfunc="sum",
    ).reset_index()
    final_data.columns = [month_label(c) if isinstance(c, pd.Timestamp) else c for c in final_data.columns]
    ordered_cols = KEY_COLS + [month_label(m) for m in all_months]
    final_data = final_data[ordered_cols].sort_values(KEY_COLS).reset_index(drop=True)
    assert_unique_key(final_data, KEY_COLS, "Final data")
    assert not (set(final_data["SB CODE"]) & OLD_CODES), "Old SB codes found in final output"
    missing_new = sorted(NEW_CODES - set(final_data["SB CODE"]))
    assert not missing_new, f"Expected new codes missing from final output: {missing_new}"

    overlap = compare_overlap(rx2025, rx2026)
    monthly_market = build_monthly_market_table(final_long)
    audit = build_transformation_audit()
    validation_summary = build_validation_summary(
        rx2025, rx2026, final_data, final_long, checks, overlap, market_check, warnings
    )
    export_excel(OUTPUT_PATH, final_data, validation_summary, monthly_market, audit, checks, overlap)
    output_check = validate_output_workbook(OUTPUT_PATH)

    overlap_material = overlap[overlap["Abs_Diff"] > 1e-6]
    result = {
        "output_path": str(OUTPUT_PATH),
        "script_path": str(Path(__file__).resolve()),
        "final_rows": len(final_data),
        "final_sanibricks": final_data["SB CODE"].nunique(),
        "period": "Jan-24 to Jun-26",
        "manufacturers": sorted(final_data["Manufacturer"].unique().tolist()),
        "measures": sorted(final_data["Measure"].unique().tolist()),
        "market_conservation": "PASS",
        "albacete": "PASS",
        "barcelona": "PASS",
        "la_rioja": "PASS",
        "asturias": "PASS",
        "overlap_rows_with_differences": int(len(overlap_material)),
        "overlap_max_abs_difference": float(overlap["Abs_Diff"].max() if len(overlap) else 0),
        "warnings": warnings,
        "output_check": output_check,
    }
    print("\nBUILD RESULT")
    for key, value in result.items():
        print(f"{key}: {value}")
    return result


if __name__ == "__main__":
    build()
