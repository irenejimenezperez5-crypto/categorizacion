"""Focused visual-QA repair of six census identifiers, without rerunning allocation."""
from __future__ import annotations
import hashlib
import inspect
import json
from pathlib import Path
import re
import xml.etree.ElementTree as ET
import zipfile
import openpyxl
from PIL import Image,ImageDraw,ImageFont
from categorization_core import code
import run_categorization as run

NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
ET.register_namespace('',NS)


def repair():
    root=run.ROOT
    target=root/'pharmacy/output/UNIVERSO_FARMACIAS_CATEGORIZADO_2027.xlsx'
    output=root/'pharmacy/output/categorization'
    report=json.loads((output/'F8_controls.json').read_text(encoding='utf-8'))
    assert run.sha(target)==report['xlsx_sha256'],'Output changed after generation; preserve user edits'
    replacement=target.with_suffix('.type_repair.xlsx')
    corrected=0
    with zipfile.ZipFile(target) as source,zipfile.ZipFile(replacement,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=6) as dest:
        for item in source.infolist():
            content=source.read(item.filename)
            if item.filename=='xl/worksheets/sheet10.xml':
                sheet=ET.fromstring(content)
                assert sheet.find(f'.//{{{NS}}}c[@r="B4"]//{{{NS}}}t').text=='sscc origen'
                for row in range(5,11):
                    cell=sheet.find(f'.//{{{NS}}}c[@r="B{row}"]')
                    peer=sheet.find(f'.//{{{NS}}}c[@r="C{row}"]')
                    raw=cell.find(f'{{{NS}}}v').text
                    sid=code(raw,10)
                    assert sid is not None and len(sid)==10
                    for child in list(cell):cell.remove(child)
                    cell.set('t','inlineStr')
                    cell.set('s',peer.get('s'))
                    ET.SubElement(ET.SubElement(cell,f'{{{NS}}}is'),f'{{{NS}}}t').text=sid
                    corrected+=1
                content=ET.tostring(sheet,encoding='utf-8')
            dest.writestr(item,content)
    assert corrected==6
    replacement.replace(target)
    book=openpyxl.load_workbook(target,read_only=True,data_only=True)
    ws=book['Recorridos pie']
    values=list(ws.iter_rows(min_row=4,max_row=10,max_col=8,values_only=True))
    assert all(isinstance(row[1],str) and len(row[1])==10 for row in values[1:])
    assert sum(book['LEEME'].iter_rows(min_row=6,max_row=6,min_col=2,max_col=2,values_only=True).__next__())==report['national_market_eur']
    names=book.sheetnames
    book.close()
    folder=output/'previews'
    widths=[170,200,240,185,185,185,185,170]
    canvas=Image.new('RGB',(sum(widths)+40,570),'white')
    draw=ImageDraw.Draw(canvas)
    regular=ImageFont.truetype('C:/Windows/Fonts/calibri.ttf',16)
    bold=ImageFont.truetype('C:/Windows/Fonts/calibrib.ttf',17)
    title=ImageFont.truetype('C:/Windows/Fonts/calibrib.ttf',26)
    draw.text((20,15),'Piloto · contraste de accesibilidad a pie',fill='#133C55',font=title)
    draw.text((20,60),'SSCC conservadas como identificadores de diez digitos. Contraste de red OSM; euros sin cambios.',fill='#243746',font=regular)
    y=125
    for i,row in enumerate(values):
        height=100 if i==0 else 46
        draw.rectangle((20,y,canvas.width-20,y+height),fill='#133C55' if i==0 else('#F2F7F8' if i%2==0 else'white'))
        x=20
        for j,(v,width) in enumerate(zip(row,widths)):
            s='' if v is None else(str(v) if isinstance(v,str) else f'{v:.6f}' if 3<=j<=6 else f'{v:,.0f}')
            limit=int((width-16)/8)
            chunks=[]
            while len(s)>limit:
                cut=s.rfind(' ',0,limit)
                if cut<=0:cut=limit
                chunks.append(s[:cut]);s=s[cut:].lstrip()
            chunks.append(s)
            for n,line in enumerate(chunks[:4 if i==0 else 2]):draw.text((x+8,y+10+n*19),line,fill='white' if i==0 else'#243746',font=bold if i==0 else regular)
            x+=width
        y+=height
    draw.text((20,535),'Recorridos pie · revision de valores exportados y tipos de identificador',fill='#666666',font=regular)
    canvas.save(folder/'Recorridos_pie.png')
    thumbs=[]
    for name in names:
        im=Image.open(folder/(re.sub(r'[^A-Za-z0-9]+','_',name)+'.png')).convert('RGB')
        im.thumbnail((950,380));thumbs.append(im)
    height=sum(max(t.height for t in thumbs[i:i+2])for i in range(0,len(thumbs),2))+30*(len(thumbs)//2+1)
    montage=Image.new('RGB',(1960,height),'#DCE8ED')
    y=15
    for i in range(0,len(thumbs),2):
        montage.paste(thumbs[i],(15,y))
        if i+1<len(thumbs):montage.paste(thumbs[i+1],(995,y))
        y+=max(t.height for t in thumbs[i:i+2])+30
    montage.save(folder/'revision_todas_hojas.png')
    report.update({'xlsx_sha256':run.sha(target),'xlsx_bytes':target.stat().st_size,'visual_review':'passed;13 sheets','post_export_type_repair':{'sheet':'Recorridos pie','cells':6,'change':'SSCC -> text,width10; all other worksheet XML byte-identical'}})
    run.write_json(output/'F8_controls.json',report)
    manifest=json.loads((output/'F8_manifest.json').read_text(encoding='utf-8'))
    upstream=json.loads((output/'F7_manifest.json').read_text(encoding='utf-8'))['signature']
    implementation=inspect.getsource(run.f8)+run.sha(Path(run.__file__).with_name('export_categorization.py'))
    signature=hashlib.sha256((upstream+implementation+run.sha(Path(run.__file__).with_name('categorization_core.py'))).encode()).hexdigest()
    manifest.update({'signature':signature,'controls':report,'post_export_type_repair':'Focused six-cell identifier repair matching current export implementation'})
    run.write_json(output/'F8_manifest.json',manifest)
    print(json.dumps({'corrected_cells':corrected,'sha256':report['xlsx_sha256'],'visual_review':report['visual_review']},ensure_ascii=True),flush=True)


if __name__=='__main__':repair()
