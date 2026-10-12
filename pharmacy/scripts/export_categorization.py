"""Openpyxl export explicitly authorized by the user; cached formulas and visual QA."""
from __future__ import annotations

from collections import defaultdict
from datetime import datetime
import hashlib
import html
import json
from pathlib import Path
import re
import zipfile

import numpy as np
import pandas as pd
from PIL import Image, ImageDraw, ImageFont
from scipy import sparse
import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.workbook.properties import CalcProperties

from categorization_core import AGE_COLUMNS, code

NAVY = "133C55"
TEAL = "087E8B"
PALE = "EAF4F6"
AMBER = "FFF2CC"
INK = "243746"
GREEN = "24704A"
ORANGE = "BC6719"
WHITE = "FFFFFF"
HEADER_ROW = 4
FIRST_ROW = 5
EUR = '#,##0.00" €";[Red]-#,##0.00" €";"–"'
POP = '#,##0.0;[Red]-#,##0.0;"–"'
PCT = '0.0%;[Red]-0.0%;"–"'


def clean(v):
    if v is None: return None
    if isinstance(v,(np.bool_,bool)): return bool(v)
    if isinstance(v,np.integer): return int(v)
    if isinstance(v,np.floating): return None if not np.isfinite(v) else float(v)
    if isinstance(v,float) and not np.isfinite(v): return None
    if pd.isna(v): return None
    return v


class Writer:
    def __init__(self):
        self.book = openpyxl.Workbook(write_only=True)
        self.book.calculation = CalcProperties(calcId=191029,fullCalcOnLoad=True,forceFullCalc=True,calcMode="auto")
        self.cache = {}
        self.frames = {}
        self.specs = {}

    def table(self,name,frame,title,note,formulas=None,labels=None,widths=None,formats=None,freeze="C5",row_height=None):
        formulas,labels,widths,formats = formulas or {},labels or {},widths or {},formats or {}
        if len(frame)+HEADER_ROW>1048576:
            raise ValueError(f"Excel row limit exceeded: {name}")
        ws = self.book.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws.sheet_format.defaultRowHeight = row_height or (64 if name in {"LEEME","Fuentes"} else 30)
        ws.freeze_panes = freeze
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.sheet_properties.outlinePr.summaryRight = False
        ws.sheet_properties.tabColor = TEAL if name in {"Universo CRM","Establecimientos"} else NAVY
        ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(frame.columns))}{max(HEADER_ROW,len(frame)+HEADER_ROW)}"
        ws.print_options.horizontalCentered = False
        ws.print_title_rows = "1:4"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = "8"  # OOXML A3; WriteOnlyWorksheet lacks constants.
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.row_dimensions[1].height = 30
        ws.row_dimensions[2].height = 42
        ws.row_dimensions[4].height = 48
        for col,c in enumerate(frame.columns,1):
            width = widths.get(c, 23 if c in formats else 25)
            if any(x in c.lower() for x in ["nombre","farmacia","direccion","metodo","motivo","observacion","detalle","url","fuente"]): width = widths.get(c,42)
            ws.column_dimensions[get_column_letter(col)].width = min(70,max(14,width))
        row1 = [WriteOnlyCell(ws,value=title)]
        row1[0].font = Font(name="Calibri",size=18,bold=True,color=NAVY)
        ws.append(row1)
        row2 = [WriteOnlyCell(ws,value=note)]
        row2[0].font = Font(name="Calibri",size=10,color=INK)
        row2[0].alignment = Alignment(wrap_text=True,vertical="top")
        # Write-only worksheets do not support merges. Use the worksheet's native
        # merge list before serialization; data/filter ranges remain unmerged.
        ws.merged_cells.add(f"A1:{get_column_letter(max(2,min(8,len(frame.columns))))}1")
        ws.merged_cells.add(f"A2:{get_column_letter(max(2,min(8,len(frame.columns))))}2")
        ws.append(row2)
        ws.append([])
        header = []
        for c in frame:
            cell = WriteOnlyCell(ws,value=labels.get(c,c.replace("_"," ")))
            cell.font = Font(name="Calibri",bold=True,size=10,color=WHITE)
            cell.fill = PatternFill("solid",fgColor=NAVY)
            cell.alignment = Alignment(wrap_text=True,vertical="center",horizontal="left")
            header.append(cell)
        ws.append(header)
        cached_columns = {j:frame[c].to_numpy() for j,c in enumerate(frame.columns,1) if c in formulas}
        self.cache[name] = cached_columns
        self.frames[name] = frame
        self.specs[name] = {"labels":labels,"formats":formats,"widths":widths,"title":title,"note":note}
        body_font = Font(name="Calibri",size=10,color=INK)
        derived_font = Font(name="Calibri",size=10,color=GREEN)
        number_alignment = Alignment(horizontal="right",vertical="center")
        text_alignment = Alignment(horizontal="left",vertical="center",wrap_text=True)
        for index,values in enumerate(frame.itertuples(index=False,name=None),FIRST_ROW):
            cells = []
            for col,(field,value) in enumerate(zip(frame.columns,values),1):
                value = clean(value)
                if field in formulas:
                    value = formulas[field](index,col)
                cell = WriteOnlyCell(ws,value=value)
                # Source strings that happen to start '=' are literal, not user formulas.
                if field not in formulas and isinstance(value,str): cell.data_type="s"
                cell.font = derived_font if field in formulas else body_font
                cell.alignment = number_alignment if isinstance(clean(values[col-1]),(int,float)) and not isinstance(clean(values[col-1]),bool) else text_alignment
                cell.number_format = formats.get(field,"General")
                if name=="LEEME" and field=="explicacion_o_valor" and index==6:cell.number_format=EUR
                if index%2==0: cell.fill = PatternFill("solid",fgColor="F2F7F8")
                cells.append(cell)
            ws.append(cells)
        for field in ["provisional","sensible","revision_identidad","ubicacion_provisional","territorio_provisional","demanda_territorial_imputada"]:
            if field in frame:
                col=get_column_letter(frame.columns.get_loc(field)+1)
                ws.conditional_formatting.add(f"{col}5:{col}{len(frame)+4}",CellIsRule(operator="equal",formula=["TRUE"],fill=PatternFill("solid",fgColor=AMBER)))
        for field in ["categoria_final","categoria_base"]:
            if field in frame:
                col = get_column_letter(frame.columns.get_loc(field)+1)
                for category,color in [("A","C6E0B4"),("B","DDEBF7"),("C","FFF2CC"),("D","F8CBAD")]:
                    ws.conditional_formatting.add(f"{col}5:{col}{len(frame)+4}",CellIsRule(operator="equal",formula=[f'"{category}"'],fill=PatternFill("solid",fgColor=color)))
        return {c:get_column_letter(j) for j,c in enumerate(frame.columns,1)}

    def save(self,target):
        temporary = target.with_suffix(".uncached.xlsx")
        self.book.save(temporary)
        # Openpyxl cannot write formula caches. Inject the independently computed,
        # validated values into standard OOXML <v> cells; formulas remain editable.
        cell_pattern = re.compile(r'(<c\b[^>]*\br="([A-Z]+)(\d+)"[^>]*>)(<f\b[^>]*>.*?</f>)(?:<v(?:\s*/>|>.*?</v>))?(</c>)')
        with zipfile.ZipFile(temporary) as source,zipfile.ZipFile(target,"w",compression=zipfile.ZIP_DEFLATED,compresslevel=6) as dest:
            sheet_number=0
            for item in source.infolist():
                data=source.read(item.filename)
                if re.fullmatch(r"xl/worksheets/sheet\d+\.xml",item.filename):
                    sheet_number=int(re.search(r"sheet(\d+)",item.filename).group(1))
                    name=self.book.sheetnames[sheet_number-1]
                    caches=self.cache[name]
                    def replace(match):
                        start,column,row,formula,end=match.groups()
                        col=column_index_from_string(column)
                        position=int(row)-FIRST_ROW
                        if col not in caches or position<0 or position>=len(caches[col]):
                            raise RuntimeError(f"Formula without validated cache {name}!{column}{row}")
                        value=clean(caches[col][position])
                        start=re.sub(r'\s+t="[^"]*"','',start)
                        if isinstance(value,str):
                            start=start[:-1]+' t="str">'
                            cache=html.escape(value)
                        elif isinstance(value,bool):
                            start=start[:-1]+' t="b">'
                            cache="1" if value else "0"
                        elif value is None:
                            cache=""
                        else:
                            cache=str(value)
                        return start+formula+"<v>"+cache+"</v>"+end
                    data=cell_pattern.sub(replace,data.decode("utf-8")).encode("utf-8")
                dest.writestr(item,data)
        temporary.unlink()  # Exact file authored here, no recursive or computed deletion.


def numeric_formats(frame):
    result={}
    for c in frame:
        if "eur" in c or c in {"Mercado","Presupuesto"}: result[c]=EUR
        elif "pct" in c or "porcentaje" in c or "variacion" in c: result[c]=PCT
        elif c.startswith("pop_") or "poblacion" in c: result[c]=POP
        elif "lon" in c or "lat" in c: result[c]="0.000000"
        elif pd.api.types.is_numeric_dtype(frame[c]) and not pd.api.types.is_bool_dtype(frame[c]): result[c]="#,##0"
    return result


def create_previews(target,writer,folder):
    """Inspect actual exported cached values, not a separate presentation data set."""
    folder.mkdir(parents=True,exist_ok=True)
    book=openpyxl.load_workbook(target,read_only=True,data_only=True)
    font_path=Path("C:/Windows/Fonts/calibri.ttf")
    bold_path=Path("C:/Windows/Fonts/calibrib.ttf")
    regular=ImageFont.truetype(str(font_path),16) if font_path.exists() else ImageFont.load_default()
    bold=ImageFont.truetype(str(bold_path),17) if bold_path.exists() else ImageFont.load_default()
    titlefont=ImageFont.truetype(str(bold_path),26) if bold_path.exists() else bold
    previews=[]
    for name,ws in zip(book.sheetnames,book):
        if name=="Universo CRM":
            columns=["crm_id","account_name","categoria_final","potencial_mercado_eur","pct_potencial_asignado","pop_65_mas","pop_75_mas","pop_85_mas"]
            positions=[writer.frames[name].columns.get_loc(c) for c in columns]
            allrows=list(ws.iter_rows(min_row=4,max_row=10,max_col=max(positions)+1,values_only=True))
            data=[tuple(row[j] for j in positions) for row in allrows]
        else:
            data=list(ws.iter_rows(min_row=4,max_row=10,max_col=min(ws.max_column or len(writer.frames[name].columns),8),values_only=True))
            columns=list(writer.frames[name].columns)[:8]
        # Summary cards/doc sheets are wider than the extensive data tables.
        widths=[min(360, max(145,writer.specs[name]["widths"].get(c,23)*7)) for c in columns]
        if name in {"LEEME","Parametros","Fuentes","Controles"}: widths=[min(500,w+100) for w in widths]
        width=sum(widths)+40
        canvas=Image.new("RGB",(int(width),570),"white")
        draw=ImageDraw.Draw(canvas)
        draw.text((20,15),writer.specs[name]["title"],fill="#"+NAVY,font=titlefont)
        note=writer.specs[name]["note"]
        maxchars=int((width-40)/8)
        lines=[]
        for words in note.split("\n"):
            while len(words)>maxchars:
                cut=words.rfind(" ",0,maxchars)
                if cut<=0: cut=maxchars
                lines.append(words[:cut]);words=words[cut:].lstrip()
            lines.append(words)
        for j,line in enumerate(lines[:3]): draw.text((20,57+j*19),line,fill="#"+INK,font=regular)
        y=125
        for rownum,row in enumerate(data):
            height=100 if rownum==0 else 46
            draw.rectangle((20,y,width-20,y+height),fill="#"+NAVY if rownum==0 else ("#F2F7F8" if rownum%2==0 else "white"))
            x=20
            for j,(value,colwidth) in enumerate(zip(row,widths)):
                if value is None: s=""
                elif isinstance(value,bool):s="Sí" if value else "No"
                elif isinstance(value,(int,float)):
                    f=writer.specs[name]["formats"].get(columns[j],"")
                    if f==EUR:s=f"{value:,.2f} €"
                    elif f==PCT:s=f"{value:.1%}"
                    elif f==POP:s=f"{value:,.1f}"
                    elif "0.000000" in f:s=f"{value:.6f}"
                    else:s=f"{value:,.0f}"
                else:s=str(value)
                maxc=max(14,int((colwidth-16)/8))
                wrapped=[]
                while len(s)>maxc and len(wrapped)<(3 if rownum==0 else 1):
                    cut=s.rfind(" ",0,maxc)
                    if cut<=0:cut=maxc
                    wrapped.append(s[:cut]);s=s[cut:].lstrip()
                wrapped.append(s[:maxc])
                for line_index,line in enumerate(wrapped[:4 if rownum==0 else 2]):
                    draw.text((x+8,y+10+line_index*19),line,fill="white" if rownum==0 else "#"+INK,font=bold if rownum==0 else regular)
                x+=colwidth
            y+=height
        draw.text((20,535),f"{name} · vista de control de valores exportados · filtros y paneles inmovilizados",fill="#666666",font=regular)
        file=folder/(re.sub(r"[^A-Za-z0-9]+","_",name)+".png")
        canvas.save(file)
        previews.append(file)
    book.close()
    thumbs=[]
    for file in previews:
        im=Image.open(file).convert("RGB");im.thumbnail((950,380));thumbs.append(im)
    height=sum(max(t.height for t in thumbs[i:i+2]) for i in range(0,len(thumbs),2))+30*(len(thumbs)//2+1)
    montage=Image.new("RGB",(1960,height),"#DCE8ED")
    y=15
    for i in range(0,len(thumbs),2):
        montage.paste(thumbs[i],(15,y))
        if i+1<len(thumbs):montage.paste(thumbs[i+1],(995,y))
        y+=max(t.height for t in thumbs[i:i+2])+30
    montage.save(folder/"revision_todas_hojas.png")
    return previews


def export_and_validate(root,config,universe,market,geo,categorized,base,sensitivity):
    output=root/"pharmacy/output/categorization"
    target=root/"pharmacy/output/UNIVERSO_FARMACIAS_CATEGORIZADO_2027.xlsx"
    stores=categorized["stores"].copy()
    sections=geo["sections"]
    budgets=market["budgets"]
    shares=sparse.load_npz(root/"pharmacy/cache/categorization/shares_base.npz")
    total_raw=np.asarray(shares.T@base["section_budget_cents"]).ravel()
    for field,mask in [("pct_potencial_sscc_imputadas",sections.provisional_section.values),("pct_potencial_respaldo_rural",sections.index.isin(base["rural_fallback_sections"]))]:
        numerator=np.asarray(shares.T@(base["section_budget_cents"]*mask)).ravel()
        stores[field]=np.divide(numerator,total_raw,out=np.zeros(len(stores)),where=total_raw>0)
    stores["demanda_territorial_imputada"]=stores.pct_potencial_sscc_imputadas.gt(1e-10)
    unmapped=geo["raw_geometry"].loc[~geo["raw_geometry"].sscc.isin(sections.sscc),["sscc","municipio","district","NPRO","NMUN"]].copy()
    unmapped["tratamiento"]="No es una unidad adicional IQVIA2026; patrones municipales para imputaciones y reescalado por SB. No anadir poblacion al total IQVIA."
    unmapped.to_csv(output/"sscc_ine2025_sin_enlace_exacto_2026.csv",index=False,encoding="utf-8-sig")
    writer=Writer()
    notes=pd.DataFrame([
        ("Entrega", "Universo CRM completo con categorias y atribucion anual de mercado RX"),
        ("Mercado nacional EUR",int(budgets.sum())/100),
        ("Filas CRM",len(universe["crm"])),
        ("Establecimientos unicos",len(stores)),
        ("Periodo", "ESSITY + RESTO; julio 2025 - junio 2026; sin proyeccion de crecimiento 2027"),
        ("Poblacion", "Tramos IQVIA por SB distribuidos con INE2025. 85+ incluye 85-89,90-94,95-99,100+. Cantidades fraccionarias atribuidas, no clientes observados."),
        ("Potencial", "Atribucion de sell-in RX de mayoristas segun documentacion IQVIA. No venta observada ni ingreso neto de cada establecimiento."),
        ("Porcentaje principal", "Potencial del establecimiento / mercado captado por todos los establecimientos situados en su SB. NO usar el RX original de ese SB como denominador."),
        ("Flujos SB", "Cada fila representa SB origen -> establecimiento. El origen conserva todo su presupuesto a centimos; el destino puede captar de vecinos."),
        ("Evitar duplicar totales", "Universo CRM conserva repetidos. Usar Establecimientos para totales o mercado_sumable_crm_eur con contabilizar_modelo=1."),
        ("Ubicacion provisional", "Posiciones precisas separadas de proxies. Hasta 16 puntos de masa en SSCC, municipio o SB; la farmacia compite UNA vez."),
        ("Ruralidad", "DEGURBA municipal 2021: ciudad e intermedia por distancia; rural por demanda municipal y reparto neutral. Respaldo SB cuando no hay puntos municipales."),
        ("Malla", "250m recortada por SSCC; densidad 65+ de malla Eurostat2021 1km, uniforme dentro de cada celda 1km. Es una aproximacion, no un censo de mayores a 250m."),
        ("Categorias", "Por euros acumulados anteriores dentro de CCAA: A<20%, B<50%, C<75%, D resto. Minimo C en SB>=P75; si todo A, menor potencial pasa a B."),
        ("Sensibilidad", "Centroide unico, exponentes1.5/3, edades1/1/1 y1/2/4. Marcar si categoria cambia o variacion EUR>20%."),
        ("Ajustes manuales", "Revisar por SB de origen en Flujos SB; redistribuir pesos normalizados y conservar presupuesto. Las formulas de resumen permanecen editables; cambiar supuestos espaciales requiere reejecutar el modelo."),
        ("Lectura inicial", "Universo CRM: localizar cuenta. Establecimientos: potencial/poblacion. Sanibricks y Controles: conciliacion. Piloto y Recorridos pie: revision."),
        ("Limitacion comercial", "Conservacion contable no prueba exactitud comercial. Revisar puntos provisionales/sensibles y contrastar con ventas por establecimiento o campo."),
    ],columns=["concepto","explicacion_o_valor"])
    def summary_formula(r,c):
        return {6:f"=SUM('Establecimientos'!E5:E{len(stores)+4})",7:f"=COUNTA('Universo CRM'!A5:A{len(universe['crm'])+4})",8:f"=COUNTA('Establecimientos'!A5:A{len(stores)+4})"}.get(r,clean(notes.explicacion_o_valor.iloc[r-FIRST_ROW]))
    writer.table("LEEME",notes,"Categorizacion de farmacias · 2027","Modelo estimado con fuentes y conciliacion por origen. Leer antes de sumar o ajustar.",formulas={"explicacion_o_valor":summary_formula},widths={"concepto":34,"explicacion_o_valor":115},freeze="A5")
    parameters=pd.DataFrame([
        ("Peso 65-74",config["age_weights"][0],"Supuesto del modelo; tramos excluyentes"),
        ("Peso 75-84",config["age_weights"][1],"Supuesto del modelo; tramos excluyentes"),
        ("Peso 85+",config["age_weights"][2],"Incluye todos los quinquenios >=85"),
        ("Exponente distancia",config["distance_exponent"],"Kernel 1/(distancia+offset)^exponente"),
        ("Offset metros",config["distance_offset_m"],"Evita singularidad a distancia cero"),
        ("Radio candidatos metros",config["candidate_radius_m"],"Ademas de SB propio y colindantes"),
        ("Malla demanda metros",config["demand_grid_m"],"Patron censal oficial2021 de1km desagregado por superficie"),
        ("Percentil SB alto",config["high_sb_quantile"],"Presupuesto RX origen nacional"),
        ("Corte A",config["category_cutoffs"][0],"Acumulado EUR anterior por CCAA"),
        ("Corte B",config["category_cutoffs"][1],"Acumulado EUR anterior por CCAA"),
        ("Corte C",config["category_cutoffs"][2],"Acumulado EUR anterior por CCAA"),
        ("Sensibilidad relativa",config["sensitivity_relative_threshold"],"Ademas de cambios de categoria"),
        ("Demografia INE",config["population_year"],"Reescalada a totales IQVIA SB2026 por edad"),
        ("DEGURBA",config["degurba_year"],"Clasificacion municipal oficial, tres clases"),
        ("Version modelo",config["model_version"],"Configuracion y hashes registrados en manifiestos"),
        ("Inicio MAT",config["period_start"],"Mes inicial incluido"),
        ("Fin MAT",config["period_end"],"Mes final incluido"),
    ],columns=["parametro","valor","explicacion"])
    # Establishments first in calculation dependencies, with straightforward formula derivations.
    fields=["establishment_id","account_name","account_type","categoria_final","potencial_mercado_eur","pct_potencial_asignado","pop_65_mas","pop_75_mas","pop_85_mas",*AGE_COLUMNS[:-1],"poblacion_ponderada","mercado_desde_sb_propio_eur","mercado_desde_vecinos_eur","mercado_rx_original_sb_eur","mercado_captado_sb_ubicacion_eur","categoria_base","sb_alto","motivo_ajuste_categoria","ccaa_modelo","sb_modelo","sscc_modelo","municipio_modelo","ruralidad","degurba","metodo_ubicacion","provisional","ubicacion_provisional","territorio_provisional","sensible","max_variacion_relativa","latitude_validada","longitude_validada","latitude_proxy","longitude_proxy","n_ubicaciones_modeladas","n_filas_crm","revision_identidad","coordenada_compartida","sb_crm","sb_iqvia","sscc_iqvia","join_method","conflicto_sb_crm_iqvia","sb_modificado_vs_crm","incidencia_ubicacion","pct_acumulado_previo_ccaa","periodo_mercado"]
    table=stores[fields].copy()
    for c in ["pct_potencial_sscc_imputadas","pct_potencial_respaldo_rural","demanda_territorial_imputada"]:table[c]=stores[c]
    fields=list(table.columns)
    letters={c:get_column_letter(j) for j,c in enumerate(fields,1)}
    formulas={
        "potencial_mercado_eur":lambda r,c:f"=SUM({letters['mercado_desde_sb_propio_eur']}{r},{letters['mercado_desde_vecinos_eur']}{r})",
        "pct_potencial_asignado":lambda r,c:f"=IFERROR({letters['potencial_mercado_eur']}{r}/{letters['mercado_captado_sb_ubicacion_eur']}{r},0)",
        "pop_65_mas":lambda r,c:"=SUM("+",".join(f"{letters[k]}{r}" for k in AGE_COLUMNS)+")",
        "pop_75_mas":lambda r,c:"=SUM("+",".join(f"{letters[k]}{r}" for k in AGE_COLUMNS[2:])+")",
        "poblacion_ponderada":lambda r,c:f"=SUM({letters[AGE_COLUMNS[0]]}{r},{letters[AGE_COLUMNS[1]]}{r})*'Parametros'!$B$5+SUM({letters[AGE_COLUMNS[2]]}{r},{letters[AGE_COLUMNS[3]]}{r})*'Parametros'!$B$6+{letters[AGE_COLUMNS[4]]}{r}*'Parametros'!$B$7"
    }
    writer.table("Establecimientos",table,"Establecimientos unicos · potencial y justificacion","Sumar esta hoja para evitar duplicados CRM. Euros y poblacion son atribuciones estimadas; proxies y sensibilidad visibles.",formulas=formulas,labels={"account_name":"Farmacia / establecimiento","account_type":"Tipo de establecimiento"},formats=numeric_formats(table),widths={"account_name":42,"establishment_id":23})
    store_rows={sid:i+FIRST_ROW for i,sid in enumerate(table.establishment_id)}
    crm=universe["crm"].copy()
    additional=["categoria_final","potencial_mercado_eur","pct_potencial_asignado","pop_65_mas","pop_75_mas",*AGE_COLUMNS,"poblacion_ponderada","mercado_desde_sb_propio_eur","mercado_desde_vecinos_eur","mercado_rx_original_sb_eur","mercado_captado_sb_ubicacion_eur","categoria_base","sb_alto","motivo_ajuste_categoria","sb_modelo","sscc_modelo","municipio_modelo","ccaa_modelo","ruralidad","metodo_ubicacion","provisional","sensible","max_variacion_relativa","latitude_validada","longitude_validada","latitude_proxy","longitude_proxy","n_filas_crm","incidencia_ubicacion"]
    additional += ["pct_potencial_sscc_imputadas","pct_potencial_respaldo_rural","demanda_territorial_imputada"]
    lookup=table.set_index("establishment_id")
    original=crm[universe["original_columns"]].copy()
    original["establishment_id"]=crm.establishment_id
    for c in additional: original[c]=crm.establishment_id.map(lookup[c])
    original["contabilizar_modelo"]=~crm.establishment_id.duplicated()
    original["mercado_sumable_crm_eur"]=original.potencial_mercado_eur*original.contabilizar_modelo
    geo_original=universe["geo_candidates"].set_index("crm_key")
    for source,field in [("latitude_final","latitud_v2_original"),("longitude_final","longitud_v2_original"),("coordinate_source_final","fuente_v2_original"),("coordinate_quality_final","calidad_v2_original")]:
        original[field]=crm.crm_key.map(geo_original[source])
    def crm_link(r,field):
        reference=f"'Establecimientos'!{letters[field]}{store_rows[crm.establishment_id.iloc[r-FIRST_ROW]]}"
        return f'=IF(ISBLANK({reference}),"",{reference})'
    crm_formulas={c:(lambda r,col,field=c:crm_link(r,field)) for c in additional}
    crm_letters={c:get_column_letter(j) for j,c in enumerate(original.columns,1)}
    crm_formulas["mercado_sumable_crm_eur"]=lambda r,c:f"={crm_letters['potencial_mercado_eur']}{r}*{crm_letters['contabilizar_modelo']}{r}"
    writer.table("Universo CRM",original,"Universo CRM · 22.356 registros conservados","Las 14 columnas originales conservan valores y orden. Repetidos enlazan al mismo establecimiento; sumar mercado_sumable_crm_eur.",formulas=crm_formulas,labels={c:c for c in universe["original_columns"]},formats=numeric_formats(original),widths={"account_name":42,"phone":22,"email":42})
    # Sparse origin-to-establishment ledger, including fractional population by origin.
    flow=base["flow_cents"].tocoo()
    order=np.lexsort((flow.col,flow.row))
    rr,cc,amount=flow.row[order],flow.col[order],flow.data[order]
    origin_sb=budgets.index.to_numpy()[rr]
    group=sparse.csr_matrix((np.ones(len(sections)),(sections.sb.map({sb:i for i,sb in enumerate(budgets.index)}),np.arange(len(sections)))),shape=(len(budgets),len(sections)))
    flows=pd.DataFrame({"sb_origen":origin_sb,"establishment_id":stores.establishment_id.to_numpy()[cc],"sb_ubicacion":stores.sb_modelo.to_numpy()[cc],"presupuesto_origen_eur":budgets.to_numpy()[rr]/100.,"pct_atribuido_origen":amount/budgets.to_numpy()[rr],"mercado_atribuido_eur":amount/100.,"ccaa_origen":market["sb_population"].ccaa.to_numpy()[rr],"ccaa_ubicacion":stores.ccaa_modelo.to_numpy()[cc],"tipo_flujo":np.where(origin_sb==stores.sb_modelo.to_numpy()[cc],"Propio","Vecino"),"mercado_sin_redondear_eur":np.asarray(base["flows_unrounded"][rr,cc]).ravel()/100.})
    for c in AGE_COLUMNS:
        matrix=(group@shares.multiply(sections[c].to_numpy()[:,None])).tocsr()
        flows[c]=np.asarray(matrix[rr,cc]).ravel()
        assert abs(flows[c].sum()-sections[c].sum())<1e-5
    flows["pop_65_mas"]=flows[AGE_COLUMNS].sum(axis=1)
    flows["pop_75_mas"]=flows[AGE_COLUMNS[2:]].sum(axis=1)
    flows["ajuste_redondeo_eur"]=flows.mercado_atribuido_eur-flows.mercado_sin_redondear_eur
    flows["contabilizar_presupuesto_origen"]=~flows.sb_origen.duplicated()
    flows["presupuesto_origen_sumable_eur"]=flows.presupuesto_origen_eur*flows.contabilizar_presupuesto_origen
    assert round(flows.presupuesto_origen_sumable_eur.sum()*100)==int(budgets.sum())
    population_origin_errors={c:float((flows.groupby("sb_origen")[c].sum()-market["sb_population"][c]).abs().max()) for c in AGE_COLUMNS}
    assert max(population_origin_errors.values())<1e-5
    fc={c:get_column_letter(j) for j,c in enumerate(flows.columns,1)}
    ff={"pct_atribuido_origen":lambda r,c:f"=IFERROR({fc['mercado_atribuido_eur']}{r}/{fc['presupuesto_origen_eur']}{r},0)","pop_65_mas":lambda r,c:"=SUM("+",".join(f"{fc[k]}{r}" for k in AGE_COLUMNS)+")","pop_75_mas":lambda r,c:"=SUM("+",".join(f"{fc[k]}{r}" for k in AGE_COLUMNS[2:])+")","ajuste_redondeo_eur":lambda r,c:f"={fc['mercado_atribuido_eur']}{r}-{fc['mercado_sin_redondear_eur']}{r}","presupuesto_origen_sumable_eur":lambda r,c:f"={fc['presupuesto_origen_eur']}{r}*{fc['contabilizar_presupuesto_origen']}{r}"}
    writer.table("Flujos SB",flows,"Flujos · Sanibrick de origen a establecimiento","Porcentaje sobre presupuesto del SB de origen. Ese presupuesto se repite por referencia; sumar solo presupuesto_origen_sumable_eur.",formulas=ff,formats=numeric_formats(flows))
    by_origin=flows.groupby("sb_origen").mercado_atribuido_eur.sum()
    captured=stores.groupby("sb_modelo").potencial_mercado_eur.sum()
    sb=pd.DataFrame({"sb":budgets.index,"ccaa":market["sb_population"].ccaa.values,"mercado_rx_origen_eur":budgets.values/100.,"mercado_atribuido_origen_eur":by_origin.reindex(budgets.index).values,"diferencia_origen_eur":np.round(by_origin.reindex(budgets.index).values-budgets.values/100.,2),"establecimientos_ubicados":stores.groupby("sb_modelo").size().reindex(budgets.index,fill_value=0).values,"mercado_captado_ubicacion_eur":captured.reindex(budgets.index,fill_value=0).values,"sb_alto":budgets.values>=budgets.quantile(config["high_sb_quantile"]),"clases_degurba":sections.groupby("sb").degurba.nunique().reindex(budgets.index).values})
    sb["diferencia_captacion_eur"]=sb.mercado_captado_ubicacion_eur-sb.mercado_rx_origen_eur
    for c in AGE_COLUMNS:sb[c]=market["sb_population"][c].values
    sb["pop_65_mas"]=sb[AGE_COLUMNS].sum(axis=1)
    sb["pop_75_mas"]=sb[AGE_COLUMNS[2:]].sum(axis=1)
    sl={c:get_column_letter(j) for j,c in enumerate(sb.columns,1)}
    sf={"diferencia_origen_eur":lambda r,c:f"=ROUND({sl['mercado_atribuido_origen_eur']}{r}-{sl['mercado_rx_origen_eur']}{r},2)","diferencia_captacion_eur":lambda r,c:f"={sl['mercado_captado_ubicacion_eur']}{r}-{sl['mercado_rx_origen_eur']}{r}","pop_65_mas":lambda r,c:"=SUM("+",".join(f"{sl[k]}{r}" for k in AGE_COLUMNS)+")","pop_75_mas":lambda r,c:"=SUM("+",".join(f"{sl[k]}{r}" for k in AGE_COLUMNS[2:])+")"}
    writer.table("Sanibricks",sb,"Sanibricks · presupuesto de origen y captacion","El presupuesto de origen se conserva. La diferencia de captacion territorial refleja los flujos entre SB; no es un error contable.",formulas=sf,formats=numeric_formats(sb))
    st=sections.drop(columns="geometry").copy()
    st["presupuesto_sscc_eur"]=base["section_budget_cents"]/100.
    st["pct_captado_sb_propio"]=base["section_own_share"]
    st["respaldo_rural_sin_puntos_municipales"]=st.index.isin(base["rural_fallback_sections"])
    st["respaldo_no_rural_sin_candidatos"]=st.index.isin(base["candidate_fallback_sections"])
    st["pop_65_mas"]=st[AGE_COLUMNS].sum(axis=1)
    st["pop_75_mas"]=st[AGE_COLUMNS[2:]].sum(axis=1)
    writer.table("SSCC",st,"Secciones censales · demanda y procedencia","Edades INE2025 reescaladas al presupuesto demografico IQVIA por SB. Geometrias y correspondencias provisionales identificadas.",formats=numeric_formats(st))
    writer.table("Sensibilidad",sensitivity["detail"],"Sensibilidad · cinco escenarios","Cambios de categoria o variaciones >20% se identifican como sensibles. Los euros nacionales se conservan en todos los escenarios.",formats=numeric_formats(sensitivity["detail"]))
    writer.table("Piloto",sensitivity["pilot"],"Piloto · revision estratificada","Semilla2027; urbano, intermedio, rural, mixto, captacion vecinal y ubicacion incierta. No es una validacion de ventas observadas.",formats=numeric_formats(sensitivity["pilot"]))
    walking=sensitivity["walking"].copy()
    walking["sscc_origen"]=walking.sscc_origen.map(lambda value:code(value,10))
    writer.table("Recorridos pie",walking,"Piloto · contraste de accesibilidad a pie","Red OSM local, vias con paso peatonal permitido. Acceso al nodo mas cercano y barreras; contraste sin cambiar el modelo geografico.",formats=numeric_formats(walking))
    distribution=categorized["category_distribution"]
    writer.table("Categorias CCAA",distribution,"Categorias · distribucion efectiva por CCAA","Version base y final: los ajustes de SB altos cambian categorias, manteniendo los mismos euros por establecimiento.",formats=numeric_formats(distribution))
    controls=[("CRM filas",len(crm),22356,"OK","Universo original completo"),("CRM ID unicos",crm.crm_id.nunique(),22356,"OK","Identificadores originales"),("Establecimientos unicos",len(stores),len(stores),"OK","39 grupos IQVIA repetidos; revisar identidad marcada"),("Mercado nacional EUR",stores.potencial_mercado_eur.sum(),budgets.sum()/100.,"OK","Conservacion a centimos"),("Max error SB origen centimos",int(np.abs(np.asarray(base['flow_cents'].sum(axis=1)).ravel()-budgets.values).max()),0,"OK","2576 presupuestos origen conservados"),("Filas duplicadas adicionales",len(crm)-len(stores),40,"OK","No compiten dos veces"),("SB altos con D",int((stores.sb_alto & stores.categoria_final.eq('D')).sum()),0,"OK","Minimo C"),("SB altos integramente A",int(stores[stores.sb_alto].groupby('sb_modelo').categoria_final.apply(lambda g:g.eq('A').all()).sum()),0,"OK","Menor potencial pasa a B"),("Ubicaciones provisionales",int(stores.ubicacion_provisional.sum()),None,"REVISAR","Distribuciones espaciales incluidas en competencia"),("Establecimientos sensibles",int(stores.sensible.sum()),None,"REVISAR","Cinco escenarios"),("Geometrias SSCC imputadas",int(sections.geometry_method.str.startswith('IMPUTADA').sum()),80,"REVISAR","Correspondencia de ediciones2025/2026 visible"),("Validacion comercial",None,None,"PENDIENTE","Contrastar ventas por punto o revision de campo")]
    for c in AGE_COLUMNS:controls.append((f"Poblacion {c}",float(stores[c].sum()),float(market['sb_population'][c].sum()),"OK","Tramo excluyente, atribucion fraccionaria"))
    controls=pd.DataFrame(controls,columns=["control","resultado","referencia","estado","detalle"])
    writer.table("Controles",controls,"Controles · conservacion y revision","OK acredita consistencia del calculo; REVISAR identifica incertidumbre. La validacion comercial requiere evidencia por establecimiento.",widths={"control":42,"detalle":90,"resultado":27,"referencia":27,"estado":18},formats={"resultado":"#,##0.00","referencia":"#,##0.00"},freeze="A5")
    writer.table("Parametros",parameters,"Parametros · supuestos y versiones","Pesos de edad en B5:B7 alimentan poblacion ponderada. Para recalcular distancias/categorias ejecutar el pipeline con la configuracion vigente.",widths={"parametro":40,"valor":25,"explicacion":90},freeze="A5")
    manifest=json.loads((output/"source_manifest.json").read_text(encoding="utf-8"))
    sources=[]
    for r in manifest['private_local_sources']:sources.append({"fuente":r['file'],"uso":"Entrada local; uso interno","referencia":r['file'],"sha256":r['sha256'],"bytes":r['bytes'],"version":"2026/entrada original"})
    for r in manifest['public_sources']['sources']:sources.append({"fuente":r['file'],"uso":"Fuente oficial; transformacion documentada","referencia":r['url'],"sha256":r['sha256'],"bytes":r['bytes'],"version":"INE2025 / DEGURBA2021 / CensusGrid2021 V3"})
    address_source=json.loads((root/"pharmacy/cache/categorization/osm_pharmacy_full_addresses_manifest.json").read_text(encoding="utf-8"))
    sources.append({"fuente":"POI farmacia con direccion completa","uso":"Recuperar portales omitidos por cache anterior","referencia":"pharmacy/cache/categorization/osm_pharmacy_full_addresses.csv.gz","sha256":address_source['extract_sha256'],"version":"Extraido del PBF local identificado por SHA256"})
    sources.extend([{"fuente":"Huff","uso":"Referencia conceptual; coeficientes propios","referencia":"https://doc.esri.com/en/arcgis-pro/latest/tool-reference/business-analyst/huff-model.html","version":"Consultado2026-10-06"},{"fuente":"INE tabla65034","uso":"Poblacion quinquenal por SSCC","referencia":"https://www.ine.es/jaxiT3/dlgExport.htm?t=65034&L=0","version":"2025"},{"fuente":"Eurostat poblacion en malla","uso":"Densidad 65+1km; subdividida250m por superficie","referencia":"https://ec.europa.eu/eurostat/en/web/gisco/geodata/population-distribution/population-grids","version":"Census2021 V3; publicacion2026"},{"fuente":"OSM licencia","uso":"POI y red peatonal local","referencia":"https://www.openstreetmap.org/copyright","version":"PBF cache local2026-09"}])
    sources=pd.DataFrame(sources)
    writer.table("Fuentes",sources,"Fuentes · versiones y huellas","Los datos IQVIA y CRM permanecen locales. Eurostat/INE se citan como fuentes oficiales; la atribucion espacial es una transformacion propia.",widths={"fuente":70,"referencia":115,"sha256":70,"uso":65},freeze="A5")
    # Reorder for the business-facing reading sequence, without breaking sheet-to-cache mapping.
    sequence=["LEEME","Universo CRM","Establecimientos","Flujos SB","Sanibricks","SSCC","Categorias CCAA","Sensibilidad","Piloto","Recorridos pie","Controles","Parametros","Fuentes"]
    writer.book._sheets=[writer.book[name] for name in sequence]
    print(f"F8 exporting {len(flows):,} origin-store flows; {len(writer.book.sheetnames)} sheets",flush=True)
    writer.save(target)
    result=validate(target,universe,stores,budgets,writer)
    previews=create_previews(target,writer,output/"previews")
    result["preview_sheets"]=len(previews)
    result["xlsx_bytes"]=target.stat().st_size
    result["xlsx_sha256"]=hashlib.sha256(target.read_bytes()).hexdigest()
    result["flow_rows"]=len(flows)
    result["population_origin_sb_max_error"]=population_origin_errors
    (output/"F8_controls.json").write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8")
    print(f"F8 {result}",flush=True)
    return result


def validate(target,universe,stores,budgets,writer):
    book=openpyxl.load_workbook(target,read_only=True,data_only=True)
    errors=[]
    counts={}
    for ws in book:
        expected=writer.frames[ws.title]
        counts[ws.title]=len(expected)
        observed=0
        for row in ws.iter_rows(min_row=FIRST_ROW,values_only=True):
            observed+=1
            # Missing-ID literals from CRM are not computed spreadsheet errors.
            computed_indices=[expected.columns.get_loc(c)-0 for c in expected.columns if c in {"potencial_mercado_eur","pct_potencial_asignado","pop_65_mas","pop_75_mas","poblacion_ponderada"}]
            if any(isinstance(row[j],str) and row[j] in {"#REF!","#DIV/0!","#VALUE!","#NAME?","#N/A","#NUM!"} for j in computed_indices):
                errors.append(f"{ws.title}:{observed+4}")
        assert observed==len(expected),f"Row count mismatch {ws.title}: {observed}/{len(expected)}"
    assert not errors,f"Formula errors: {errors[:10]}"
    ws=book['Universo CRM']
    crm=universe['crm'][universe['original_columns']]
    indices={c:j for j,c in enumerate(writer.frames['Universo CRM'].columns)}
    unique_total=0.
    checked=0
    for original,row in zip(crm.itertuples(index=False,name=None),ws.iter_rows(min_row=FIRST_ROW,values_only=True)):
        for a,b in zip(original,row[:len(original)]):
            assert clean(a)==clean(b),f"Original CRM value changed, row{checked+1}"
        unique_total+=row[indices['mercado_sumable_crm_eur']]
        checked+=1
    assert checked==22356
    assert round(unique_total*100)==int(budgets.sum())
    data=book['Establecimientos']
    columns={c:j for j,c in enumerate(writer.frames['Establecimientos'].columns)}
    total=0.
    for expected,row in zip(stores.itertuples(),data.iter_rows(min_row=FIRST_ROW,values_only=True)):
        assert row[columns['establishment_id']]==expected.establishment_id
        assert row[columns['categoria_final']]==expected.categoria_final
        total+=row[columns['potencial_mercado_eur']]
        assert abs(row[columns['pop_65_mas']]-expected.pop_65_mas)<1e-6
    assert round(total*100)==int(budgets.sum())
    book.close()
    # Inspect styles and formula structures without rereading the entire giant workbook.
    formula_count=0
    with zipfile.ZipFile(target) as z:
        assert z.testzip() is None
        for name in z.namelist():
            if re.fullmatch(r"xl/worksheets/sheet\d+\.xml",name):
                raw=z.read(name)
                formula_count+=raw.count(b'<f>')
                assert b'<autoFilter ' in raw and b'<pane ' in raw
                assert b'showGridLines="0"' in raw
                assert b'#REF!' not in raw
    assert formula_count>100000
    return {"state":"complete","sheets":counts,"crm_original_cells_verified":checked*len(universe['original_columns']),"cached_formulas_verified":formula_count,"formula_error_count":len(errors),"national_market_eur":int(budgets.sum())/100.,"national_difference_cents":0,"original_rows":checked,"physical_establishments":len(stores)}
