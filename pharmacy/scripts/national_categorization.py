"""Current V2 delivery: calibrated EUR, independent national/CCAA grades and SB priority.

Reuses the signed V1 spatial population; never reruns geocoding or alters raw inputs.
Run: pharmacy/.venv/Scripts/python.exe pharmacy/scripts/national_categorization.py
"""
from __future__ import annotations
import argparse
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
import json
from pathlib import Path
import pickle
import shutil
import sys

import numpy as np
import pandas as pd
import openpyxl
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from economic_breaks import natural_breaks, assign
from rural_allocation import rural_allocation_weights
from review_categorization import cents_allocation, plain
from export_categorization import Writer, create_previews, numeric_formats, EUR, POP, PCT

ROOT = Path(__file__).resolve().parents[2]
CACHE = ROOT / 'pharmacy/cache/categorization'
OUT = ROOT / 'pharmacy/output/current'
ARCHIVE = ROOT / 'pharmacy/output/archive/2027.1'
APPDATA = ROOT / 'visualizacion/data'
BOOKNAME = 'UNIVERSO_FARMACIAS_CATEGORIZADO_2027.xlsx'
AGES = ['pop_65_69','pop_70_74','pop_75_79','pop_80_84','pop_85_mas']
VERSION = '2027.3'
sys.path.insert(0,str(ROOT / 'visualizacion/scripts'))
from prepare_data import sha256, open_source_read


def json_file(path, value):
    path.write_text(json.dumps(plain(value),ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')


def archive_v1():
    ARCHIVE.mkdir(parents=True,exist_ok=True)
    original = ROOT / 'pharmacy/output' / BOOKNAME
    archived = ARCHIVE / BOOKNAME
    if original.exists() and not archived.exists():
        before=sha256(original)
        with open_source_read(original) as source, archived.open('wb') as destination:
            shutil.copyfileobj(source,destination)
        assert sha256(archived)==before
    if original.exists():
        assert sha256(original)==sha256(archived)
        try:
            original.unlink() # Exact historical output, only after its verified archive exists.
        except PermissionError:
            print('Excel V1 abierto: se conserva también en la ruta antigua. La entrega vigente se crea en output/current.',flush=True)
    for name in ['stores.json','metadata.json','establecimientos.csv','universo_crm.csv','flows.json','export_controls.json']:
        target=ARCHIVE/name
        if not target.exists():
            source=APPDATA/name
            shutil.copy2(source,target)
            assert sha256(source)==sha256(target)
    if not (ARCHIVE/'manifest.json').exists():
        json_file(ARCHIVE/'manifest.json',{'version':'2027.1','archived_at':datetime.now(timezone.utc).isoformat(),
                  'sha256':{p.name:sha256(p) for p in ARCHIVE.iterdir() if p.is_file()},
                  'note':'Entrega histórica, flujos y categorías acumuladas V1; no mezclar con euros calibrados V2.'})
    return ARCHIVE/'stores.json'


def sb_priority(points):
    """Policy: both national upper-quartile A/B count and A/B EUR; no local quotas."""
    frame=points.groupby('sb_code',sort=True).agg(sb_name=('sb_name','first'),ccaa=('ccaa','first'),
        province=('province','first'),establishments=('id','size'),market_eur=('potential_eur','sum'),
        observed_market_eur=('sb_original_eur','first'),provisional_count=('location_provisional','sum')).reset_index()
    for letter in 'ABCD':
        g=points[points.category_national.eq(letter)].groupby('sb_code').agg(n=('id','size'),eur=('potential_eur','sum'))
        frame[f'national_{letter}_n']=frame.sb_code.map(g.n).fillna(0).astype(int)
        frame[f'national_{letter}_eur']=frame.sb_code.map(g.eur).fillna(0.)
    frame['national_AB_n']=frame.national_A_n+frame.national_B_n
    frame['national_AB_eur']=np.round(frame.national_A_eur+frame.national_B_eur,2)
    frame['national_AB_market_pct']=frame.national_AB_eur/frame.market_eur
    eligible=points[points.visit_eligible]
    frame['visitable_establishments']=frame.sb_code.map(eligible.groupby('sb_code').size()).fillna(0).astype(int)
    frame['rural_establishments']=frame.sb_code.map(points[points.geographic_scope.eq('Rural')].groupby('sb_code').size()).fillna(0).astype(int)
    frame['unresolved_establishments']=frame.establishments-frame.visitable_establishments-frame.rural_establishments
    for letter in 'AB':
        g=eligible[eligible.category_national.eq(letter)].groupby('sb_code').agg(n=('id','size'),eur=('potential_eur','sum'))
        frame[f'visitable_national_{letter}_n']=frame.sb_code.map(g.n).fillna(0).astype(int)
        frame[f'visitable_national_{letter}_eur']=frame.sb_code.map(g.eur).fillna(0.)
    frame['visitable_national_AB_n']=frame.visitable_national_A_n+frame.visitable_national_B_n
    frame['visitable_national_AB_eur']=np.round(frame.visitable_national_A_eur+frame.visitable_national_B_eur,2)
    frame['visitable_national_AB_market_pct']=frame.visitable_national_AB_eur/frame.market_eur
    count_quantile=float(frame.visitable_national_AB_n.quantile(.75,interpolation='linear'))
    count_cut=max(1,int(np.ceil(count_quantile)))
    money_quantile=float(frame.visitable_national_AB_eur.quantile(.75,interpolation='linear'))
    money_cut=float(np.ceil(money_quantile*100-1e-7)/100)
    frame['priority_count_pass']=frame.visitable_national_AB_n.ge(count_cut)
    frame['priority_market_pass']=frame.visitable_national_AB_eur.ge(money_cut)
    frame['sb_priority_national']=frame.priority_count_pass & frame.priority_market_pass
    frame['sb_priority_reason']=np.select([frame.visitable_establishments.eq(0),frame.sb_priority_national,
        ~frame.priority_count_pass & ~frame.priority_market_pass,~frame.priority_count_pass],
        ['Sin establecimientos no rurales visitables','Cumple cantidad y euros A/B nacionales no rurales',
         'No alcanza cantidad ni euros A/B nacionales no rurales','No alcanza cantidad A/B nacional no rural'],
        default='No alcanza euros A/B nacionales no rurales')
    frame['priority_rank_national']=frame.visitable_national_AB_eur.rank(method='first',ascending=False).astype(int)
    config={'scope':'ESPAÑA','eligible_scope':'No rural: DEGURBA 1 y 2','reference_sb':len(frame),
            'quantile':.75,'quantile_interpolation':'linear','AB_count_quantile_raw':count_quantile,
            'AB_count_min':count_cut,'AB_market_quantile_raw_eur':money_quantile,'AB_market_min_eur':money_cut,
            'operator':'AND','priority_sb_count':int(frame.sb_priority_national.sum()),
            'definition':'Cantidad y euros de A/B nacionales NO RURALES >= P75 entre todos los SB, incluidos ceros.',
            'policy':'P75 es política de foco; rural no se visita y ámbito sin resolver requiere revisión.',
            'filters':'Filtros y categorías CCAA no recalculan la prioridad; rural nunca se hace visitable por pertenecer a un SB prioritario.'}
    return frame,config


def csv(path,frame):
    # Exact field IDs for visualization/automation; Spanish labels in the workbook/dictionary.
    frame.to_csv(path,index=False,sep=';',decimal=',',encoding='utf-8-sig')


LABELS={'id':'ID establecimiento único','name':'Nombre establecimiento','potential_eur':'Potencial anual estimado · €',
 'category_national':'Categoría nacional','category_ccaa':'Categoría CCAA','dual_category_disagreement':'Letras distintas',
 'sb_priority_national':'SB visita prioritaria · nacional','sb_priority_reason':'Motivo prioridad SB',
 'national_AB_n':'Farmacias A/B nacionales en SB','national_AB_eur':'Mercado A/B nacional · €',
 'national_AB_market_pct':'% mercado SB en A/B nacionales','share':'% mercado asignado del SB',
 'pop65':'Población atribuida 65+','pop75':'Población atribuida 75+','pop85':'Población atribuida 85+',
 'weighted_population':'Demanda demográfica ponderada','sb_weighted_population':'Demanda total puntos SB',
 'demographic_share':'% demanda ponderada SB','unrounded_potential_eur':'Potencial previo al redondeo · €',
 'rounding_adjustment_eur':'Ajuste redondeo a céntimos · €','sb_original_eur':'Mercado RX observado agregado SB · €',
 'sb_captured_eur':'Mercado atribuido total puntos SB · €','potential_v1_eur':'Potencial histórico V1 · €',
 'category_v1':'Categoría histórica V1','delta_potential_eur':'Cambio respecto V1 · €',
 'geographic_scope':'Ámbito rural / no rural', 'visit_eligible':'Elegible visita no rural',
 'allocation_weight':'Peso de reparto aplicado', 'sb_allocation_weight':'Peso total de reparto SB',
 'allocation_share':'% peso de reparto aplicado', 'allocation_method':'Método de reparto',
 'rural_group_key':'Grupo rural municipio–SB','rural_group_count':'Puntos únicos del grupo rural',
 'rural_group_demand':'Demanda total grupo rural',
 'visitable_national_AB_n':'A/B nacionales no rurales en SB', 'visitable_national_AB_eur':'Mercado A/B no rural · €',
 'visitable_national_AB_market_pct':'% mercado SB en A/B no rurales',
 'visitable_national_A_n':'A nacionales no rurales', 'visitable_national_B_n':'B nacionales no rurales',
 'visitable_national_A_eur':'Mercado A no rural · €','visitable_national_B_eur':'Mercado B no rural · €',
 'visitable_establishments':'Puntos no rurales visitables','rural_establishments':'Puntos rurales','unresolved_establishments':'Puntos ámbito sin resolver',
 'near_national_cut_5pct':'A menos del 5% de un corte nacional','near_ccaa_cut_5pct':'A menos del 5% de un corte CCAA',
 'location_provisional':'Ubicación estimada','territory_provisional':'Territorio provisional','provisional':'Resultado provisional',
 'location_method':'Método de ubicación','municipality_method':'Método de municipio','sensitive_v1':'Sensibilidad espacial histórica V1',
 'sensitivity_variation_v1':'Variación espacial histórica V1','pop_65_69':'Población atribuida 65–69',
 'pop_70_74':'Población atribuida 70–74','pop_75_79':'Población atribuida 75–79',
 'pop_80_84':'Población atribuida 80–84','pop_85_mas':'Población atribuida 85+',
 'ccaa':'CCAA','province':'Provincia','province_code':'Código provincia','municipality':'Municipio',
 'municipality_code':'Código municipio','postal_code':'Código postal','sb_code':'SB código','sb_name':'SB nombre',
 'sscc':'Sección censal','iqvia':'Código IQVIA','address':'Dirección','type':'Tipo establecimiento',
 'lat':'Latitud visualización','lon':'Longitud visualización','rurality':'DEGURBA · ciudad / intermedia / rural','degurba':'DEGURBA',
 'crm_ids':'Filas CRM vinculadas','crm_row_count':'Número filas CRM','establishment_id':'ID establecimiento único',
 'contabilizar_modelo':'Contabilizar esta fila CRM','mercado_sumable_crm_eur':'Mercado sumable CRM · €',
 'scope':'Referencia territorial','C_from_eur':'C desde · € (incluido)','B_from_eur':'B desde · € (incluido)',
 'A_from_eur':'A desde · € (incluido)','reference_establishments':'Establecimientos referencia',
 'reference_market_eur':'Mercado referencia · €','goodness_fit_pct':'Variación explicada por los grupos',
 'establishments':'Establecimientos únicos','market_eur':'Mercado atribuido · €','observed_market_eur':'Mercado RX observado agregado · €',
 'priority_rank_national':'Orden nacional por mercado A/B','provisional_count':'Ubicaciones estimadas',
 'priority_count_pass':'Cumple cantidad A/B','priority_market_pass':'Cumple euros A/B'}


def export_book(points,crm,cuts,sb,config,controls,sources):
    writer=Writer()
    notes=[('Entrega','V2 nacional · categorías comparables España + posición regional CCAA. Sustituye V1, archivada.'),
      ('Mercado de referencia','ESSITY + RESTO RX sell-in · julio 2025–junio 2026. Sin proyección de crecimiento a 2027.'),
      ('Qué representa','Potencial atribuido del mercado agregado, no venta observada por farmacia, clientes reales ni crecimiento Essity.'),
      ('Universo','22.356 filas CRM conservadas; 22.316 establecimientos físicos. Establecimientos es la hoja sumable.'),
      ('Reparto de euros','Presupuesto observado del SB × peso aplicado / suma de pesos SB. No rural conserva demanda espacial; rural usa demanda total del grupo municipio–SB / número de puntos rurales. Céntimos por mayores restos, empate por ID.'),
      ('Población y cercanía','Se conserva la captación espacial V1 por SSCC: distancia gradual en no rural y reparto municipal en rural. 65–74 ×1, 75–84 ×2, 85+ ×3.'),
      ('Ámbito y visita','DEGURBA original: 1 Ciudad y 2 Área intermedia = No rural visitable; 3 Rural = no visitar; 0 sin resolver = revisión. Usar filtros Ámbito y Elegible visita no rural en las dos hojas y CSV.'),
      ('Rural sin distancias','Mismo peso aplicado por farmacia dentro de cada municipio–SB rural. Neutraliza diferencias por captación vecina; conserva la demanda conjunta y el presupuesto SB. Si falta municipio, grupo rural provisional dentro del SB. Edades espaciales originales siguen visibles, no son el peso de reparto rural.'),
      ('Dos letras','Las dos usan exactamente los mismos euros. Nacional permite comparar todo el país; CCAA expresa posición relativa dentro de su mercado, con cortes distintos.'),
      ('Referencia de cortes','Cortes aprendidos de la entrega 2027.2 mediante cuatro grupos óptimos en euros, congelados al restaurar ruralidad para conservar comparabilidad. Límite en punto medio redondeado hacia arriba al céntimo. No se recalculan al filtrar.'),
      ('Límites inclusivos','A ≥ corte A; B ≥ corte B y < A; C ≥ corte C y < B; D < corte C. No hay cuotas ni protecciones por SB.'),
      ('Uso de categorías','Usar potencial € y letra nacional para comparar valor. Letra CCAA para contexto regional. A = potencial alto en su referencia; D = menor, no ausencia de oportunidad.'),
      ('Filtro de visita SB',f"Prioridad nacional si reúne al menos {config['AB_count_min']} A/B nacionales NO RURALES Y {config['AB_market_min_eur']:,.2f} € A/B nacionales NO RURALES. Ambos límites son P75 sobre {config['reference_sb']} SB, incluidos los ceros."),
      ('Por qué ambas condiciones','Cantidad evita priorizar solo una farmacia aislada; euros evita priorizar muchos puntos de valor conjunto reducido. P75 es política de foco, no retorno comercial demostrado ni ruta.'),
      ('Filtrar el Excel','Usar filtros de cabecera, en fila 4: CCAA, provincia, municipio, CP, SB, ambas categorías y SB visita prioritaria. SB prioritarios lista solo los seleccionados nacionalmente.'),
      ('CRM sin doble conteo','Los resultados se repiten para todas las filas CRM vinculadas. Sumar mercado_sumable_crm_eur o filtrar contabilizar_modelo=TRUE; nunca sumar potencial en todas las filas CRM.'),
      ('Porcentaje asignado','share = potencial de farmacia / mercado agregado observado del SB de ubicación. Dentro de cada SB suma 100% sobre establecimientos únicos.'),
      ('Edades','65+, 75+ y 85+ son acumulados solapados y no se suman. Los cinco tramos excluyentes son fraccionarios y justifican demanda ponderada.'),
      ('Histórico y flujos','Las categorías, euros propio/vecino y flujos de V1 son históricos. No son la descomposición del potencial calibrado V2; solo se conserva su señal espacial de población.'),
      ('Incertidumbre','Ubicaciones y correspondencias provisionales siguen visibles. Sensibilidad V1 no acredita estabilidad de letras V2. Cercanía ≤5% de un corte señala posible cambio con revisiones pequeñas.'),
      ('Actualizaciones','Referencias congeladas para comparar. Regenerar Excel, CSV y aplicación con el script. La referencia económica 2027.2 está congelada; revisar expresamente su sustitución. Cambiar una celda de cortes solo recalcula letras; resúmenes/prioridad/exportaciones son instantáneas y requieren regeneración.'),
      ('Datos locales','No publicar CRM ni IQVIA. Arrancar con arrancar_app.cmd; no requiere descargar dependencias.')]
    writer.table('LEEME',pd.DataFrame(notes,columns=['tema','explicacion_o_valor']),
      'Farmacias 2027 · guía de lectura V2','Excel y CSV · ámbito rural/no rural y visita · 7 octubre 2026',
      widths={'tema':28,'explicacion_o_valor':70},freeze='A5',row_height=94)
    writer.table('Cortes €',cuts,'España y CCAA · cortes derivados de los datos',
      'Cada referencia usa los mismos euros calibrados. Los límites inferiores C/B/A están incluidos. No son cortes de rentabilidad.',
      labels=LABELS,formats=numeric_formats(cuts),widths={'scope':34},freeze='B5')
    cutrows={r.scope:i+5 for i,r in cuts.iterrows()}
    cutcols={c:get_column_letter(i+1) for i,c in enumerate(cuts)}
    sbordered=sb.sort_values('sb_code').reset_index(drop=True)
    sbrows={r.sb_code:i+5 for i,r in sbordered.iterrows()}
    parameters=pd.DataFrame([('version',VERSION),('edad_65_74',1),('edad_75_84',2),('edad_85_mas',3),
       ('SB mínimo A/B nacionales',config['AB_count_min']),('SB mínimo € A/B nacionales',config['AB_market_min_eur']),
       ('Percentil de foco',config['quantile']),('Métrica para cortes','Euros calibrados, sin transformación logarítmica'),
       ('Número de grupos',4),('Cercanía al corte',.05)],columns=['parametro','valor'])
    writer.table('Parámetros',parameters,'Parámetros y política de foco','Pesos de edad y P75 son decisiones de modelo; documentar cambios y regenerar la entrega.',freeze='A5',widths={'parametro':37,'valor':57})
    leading=['id','name','potential_eur','category_national','category_ccaa','geographic_scope','rurality','visit_eligible','allocation_method','sb_priority_national','share','pop65','pop75','pop85',
             'ccaa','province','municipality','postal_code','sb_code','sb_name','sscc','iqvia','type','address']
    # Runtime compatibility aliases are intentionally omitted from the analytical book/CSV.
    discard=['category','category_base','sensitive','sensitivity_variation']
    pointframe=points[[*leading,*[c for c in points if c not in leading+discard]]].copy()
    pointframe['crm_ids']=pointframe.crm_ids.map(lambda v:' | '.join(map(str,v)))
    pointframe=pointframe.sort_values(['potential_eur','id'],ascending=[False,True]).reset_index(drop=True)
    cols={c:get_column_letter(i+1) for i,c in enumerate(pointframe)}
    def grade_formula(r,scope):
        n=cutrows[scope];money=f'{cols["potential_eur"]}{r}'
        return f'=IF({money}>=\'Cortes €\'!${cutcols["A_from_eur"]}${n},"A",IF({money}>=\'Cortes €\'!${cutcols["B_from_eur"]}${n},"B",IF({money}>=\'Cortes €\'!${cutcols["C_from_eur"]}${n},"C","D")))'
    formulas={
      'category_national':lambda r,c:grade_formula(r,'ESPAÑA'),
      'category_ccaa':lambda r,c:grade_formula(r,pointframe.iloc[r-5].ccaa),
      'dual_category_disagreement':lambda r,c:f'={cols["category_national"]}{r}<>{cols["category_ccaa"]}{r}',
      'weighted_population':lambda r,c:f'=SUM({cols[AGES[0]]}{r},{cols[AGES[1]]}{r})*\'Parámetros\'!$B$6+SUM({cols[AGES[2]]}{r},{cols[AGES[3]]}{r})*\'Parámetros\'!$B$7+{cols[AGES[4]]}{r}*\'Parámetros\'!$B$8',
      'demographic_share':lambda r,c:f'={cols["weighted_population"]}{r}/{cols["sb_weighted_population"]}{r}',
      'allocation_share':lambda r,c:f'={cols["allocation_weight"]}{r}/{cols["sb_allocation_weight"]}{r}',
      'unrounded_potential_eur':lambda r,c:f'={cols["sb_original_eur"]}{r}*{cols["allocation_share"]}{r}',
      'potential_eur':lambda r,c:f'=ROUND({cols["unrounded_potential_eur"]}{r},2)+{cols["rounding_adjustment_eur"]}{r}',
      'share':lambda r,c:f'={cols["potential_eur"]}{r}/{cols["sb_original_eur"]}{r}',
      'pop65':lambda r,c:'='+'+'.join(f'{cols[a]}{r}' for a in AGES),
      'pop75':lambda r,c:'='+'+'.join(f'{cols[a]}{r}' for a in AGES[2:]),
      'pop85':lambda r,c:f'={cols[AGES[4]]}{r}',
      'delta_potential_eur':lambda r,c:f'={cols["potential_eur"]}{r}-{cols["potential_v1_eur"]}{r}',
    }
    fmt=numeric_formats(pointframe)
    for c in ['pop65','pop75','pop85','weighted_population','sb_weighted_population']:fmt[c]=POP
    for c in ['share','demographic_share','allocation_share']:fmt[c]=PCT
    writer.table('Establecimientos',pointframe,'Universo · establecimientos físicos únicos',
      'Hoja sumable. Ambas categorías y prioridad nacional SB siempre visibles. Estimaciones, nunca venta observada por farmacia.',
      formulas=formulas,labels=LABELS,formats=fmt,widths={'name':44,'potential_eur':26,'category_national':20,'category_ccaa':20,'sb_priority_national':25})
    crmcols={c:get_column_letter(i+1) for i,c in enumerate(crm)}
    lookup={}
    for c in ['potential_eur','category_national','category_ccaa','sb_priority_national','share','pop65','pop75','pop85']:
        pcol=cols[c]
        lookup[c]=lambda r,j,k=pcol:f'=INDEX(\'Establecimientos\'!${k}$5:${k}${len(pointframe)+4},MATCH({crmcols["establishment_id"]}{r},\'Establecimientos\'!${cols["id"]}$5:${cols["id"]}${len(pointframe)+4},0))'
    lookup['mercado_sumable_crm_eur']=lambda r,j:f'=IF({crmcols["contabilizar_modelo"]}{r},{crmcols["potential_eur"]}{r},0)'
    writer.table('CRM completo',crm,'CRM original · todas las filas conservadas',
      'Las primeras columnas conservan el CRM original. Repeticiones físicas no suman: usar Mercado sumable CRM o Contabilizar esta fila CRM=TRUE.',
      formulas=lookup,labels=LABELS,formats={**numeric_formats(crm),'share':PCT,'pop65':POP,'pop75':POP,'pop85':POP},widths={'account_name':43})
    sbcols={c:get_column_letter(i+1) for i,c in enumerate(sbordered)}
    sbformula={'sb_priority_national':lambda r,j:f'=AND({sbcols["visitable_national_AB_n"]}{r}>=\'Parámetros\'!$B$9,{sbcols["visitable_national_AB_eur"]}{r}>=\'Parámetros\'!$B$10)',
               'national_AB_market_pct':lambda r,j:f'={sbcols["national_AB_eur"]}{r}/{sbcols["market_eur"]}{r}'}
    writer.table('Sanibricks',sbordered,'Sanibricks · mercado y concentración nacional A/B',
      'Prioridad calculada sobre la referencia nacional completa. Filtrar TRUE para los SB de visita prioritaria; categorías CCAA no intervienen.',
      formulas=sbformula,labels=LABELS,formats=numeric_formats(sbordered),widths={'sb_name':39,'sb_priority_reason':58})
    writer.table('SB prioritarios',sb[sb.sb_priority_national].sort_values('visitable_national_AB_eur',ascending=False),
      'SB de visita prioritaria · referencia nacional',f"Ambas condiciones: ≥{config['AB_count_min']} establecimientos A/B nacionales no rurales y ≥{config['AB_market_min_eur']:,.2f} € A/B no rurales. No cambia con los filtros.",
      labels=LABELS,formats=numeric_formats(sb),widths={'sb_name':38})
    distribution=[]
    for scope,g in [('ESPAÑA',points),*points.groupby('ccaa',sort=True)]:
        for ref,col in [('Nacional','category_national'),('CCAA','category_ccaa')]:
            for letter in 'ABCD':
                part=g[g[col].eq(letter)]
                distribution.append({'scope':scope,'referencia_categoria':ref,'categoria':letter,'establishments':len(part),
                                     'market_eur':round(float(part.potential_eur.sum()),2),'market_pct':float(part.potential_eur.sum()/g.potential_eur.sum())})
    dist=pd.DataFrame(distribution)
    writer.table('Distribución',dist,'Distribuciones de mercado y categorías','Instantánea de la entrega; sin cuotas impuestas. La distribución regional nacional agrega referencias CCAA diferentes.',labels=LABELS,formats=numeric_formats(dist))
    dictionary=pd.DataFrame([{'campo_csv':c,'etiqueta_excel':LABELS.get(c,c.replace('_',' ')),
                             'uso': 'Histórico V1; no usar como resultado vigente' if '_v1' in c else 'Indicador vigente; ver LEEME y Parámetros'} for c in pointframe])
    writer.table('Diccionario',dictionary,'Campos de Excel, CSV y aplicación','Identificadores estables en CSV; las etiquetas del Excel traducen su significado. Population = personas atribuidas fraccionarias.',widths={'campo_csv':37,'etiqueta_excel':48,'uso':64})
    writer.table('Fuentes',pd.DataFrame([{'fuente':k,'sha256_o_url':v} for k,v in sources.items()]),'Fuentes y trazabilidad','Entradas originales conservadas; hashes de los artefactos de origen. Datos privados locales.',widths={'fuente':55,'sha256_o_url':70},freeze='A5')
    writer.table('Controles',pd.DataFrame([{'control':k,'resultado':str(v)} for k,v in controls.items()]),'Conciliación y controles','Controles sobre establecimientos únicos y presupuestos de ubicación; población espacial V1 conservada.',widths={'control':54,'resultado':60},freeze='A5')
    for name in ['Establecimientos','CRM completo']:
        frame=writer.frames[name];ws=writer.book[name]
        for field in ['category_national','category_ccaa','category_v1']:
            col=get_column_letter(frame.columns.get_loc(field)+1)
            for letter,color in [('A','C6E0D5'),('B','DCEAF3'),('C','FFF0D2'),('D','E8DFEC')]:
                ws.conditional_formatting.add(f'{col}5:{col}{len(frame)+4}',CellIsRule(operator='equal',formula=[f'"{letter}"'],fill=PatternFill('solid',fgColor=color)))
    book=OUT/BOOKNAME
    writer.save(book)
    create_previews(book,writer,OUT/'previews')
    return pointframe,dist


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--no-excel',action='store_true',help='Only regenerate analytical tables, preserving an existing workbook.')
    args=parser.parse_args()
    OUT.mkdir(parents=True,exist_ok=True)
    source=archive_v1()
    historical=json.loads((ARCHIVE/'metadata.json').read_text(encoding='utf-8'))
    geometry_manifest=json.loads((APPDATA/'geometry_manifest.json').read_text(encoding='utf-8'))
    source_manifest=json.loads((ROOT/'pharmacy/output/categorization/source_manifest.json').read_text(encoding='utf-8'))
    config=json.loads((ROOT/'pharmacy/categorization_config.json').read_text(encoding='utf-8'))
    if config!=source_manifest['config']:
        raise ValueError('Han cambiado parámetros de la base espacial. Reconstruir y validar F1–F6 antes de sustituir su referencia congelada.')
    # Reorganization preserves bytes; old source paths remain historical provenance.
    guarded={'pharmacy/univero_pharmacy_crm.xlsx','pharmacy/rx/RX 2026.xlsx','Estructura SB_No Onekey_2026.xlsx',
             'pharmacy/output/PHARMACY_MASTER_GEOCODED_V2.xlsx','pharmacy/output/RX_2024_2026_SB2026.xlsx',
             'pharmacy/cache/categorization/universe.pickle','pharmacy/cache/categorization/market.pickle',
             'pharmacy/cache/categorization/attribution.pickle'}
    for entry in historical['sources']:
        name=entry['path'].replace('\\','/')
        if name not in guarded:continue
        path=ROOT/name
        if not path.exists() and (ROOT/'pharmacy/input/iqvia'/path.name).exists():path=ROOT/'pharmacy/input/iqvia'/path.name
        if sha256(path)!=entry['sha256']:raise ValueError('Base espacial o fuente privada cambiada: '+name)
    if sha256(CACHE/'market.pickle')!=geometry_manifest['market_sha256']:
        raise ValueError('El mercado ya no coincide con la referencia espacial revisada.')
    s=pd.DataFrame(json.loads(source.read_text(encoding='utf-8'))).set_index('id').sort_index()
    with (CACHE/'market.pickle').open('rb') as f:market=pickle.load(f)
    with (CACHE/'universe.pickle').open('rb') as f:universe=pickle.load(f)
    inputhashes={str(p.relative_to(ROOT)):sha256(p) for p in [source,CACHE/'market.pickle',CACHE/'universe.pickle',
         ROOT/'pharmacy/scripts/economic_breaks.py',ROOT/'pharmacy/scripts/rural_allocation.py',ROOT/'pharmacy/scripts/national_categorization.py',
         ROOT/'pharmacy/scripts/review_categorization.py',ROOT/'pharmacy/scripts/export_categorization.py']}
    s=s.rename(columns={'potential_eur':'potential_v1_eur','category':'category_v1','category_base':'category_base_v1',
       'category_reason':'category_reason_v1','sb_high':'sb_high_v1','own_eur':'own_eur_v1','neighbor_eur':'neighbor_eur_v1',
       'share':'share_v1','sb_captured_eur':'sb_captured_eur_v1','sensitive':'sensitive_v1','sensitivity_variation':'sensitivity_variation_v1'})
    expected=market['budgets'].reindex(s.sb_code).to_numpy()
    assert np.array_equal(np.rint(s.sb_original_eur.to_numpy()*100).astype('int64'),expected)
    # Independently extracted manufacturer/MAT table from the completed R1 audit.
    raw=ROOT/'pharmacy/analysis/categorization_review_20261006/market_manufacturer_sb.csv'
    observed=pd.read_csv(raw,sep=';',decimal=',',dtype={'sb':str}).set_index('sb')
    observed_cents=np.rint(observed[['ESSITY_Euros','RESTO_Euros']].sum(axis=1)*100).astype('int64').sort_index()
    assert observed_cents.equals(market['budgets'].sort_index())
    inputhashes[str(raw.relative_to(ROOT))]=sha256(raw)
    s=rural_allocation_weights(s)
    s['potential_eur']=cents_allocation(s.allocation_weight,market['budgets']/100,s.sb_code)/100
    s['sb_weighted_population']=s.groupby('sb_code').weighted_population.transform('sum')
    s['demographic_share']=s.weighted_population/s.sb_weighted_population
    s['unrounded_potential_eur']=s.sb_original_eur*s.allocation_share
    excel_rounded=s.unrounded_potential_eur.map(lambda v:float(Decimal(str(v)).quantize(Decimal('.01'),rounding=ROUND_HALF_UP)))
    s['rounding_adjustment_eur']=np.round(s.potential_eur-excel_rounded,2)
    s['share']=s.potential_eur/s.sb_original_eur
    s['sb_captured_eur']=s.sb_original_eur
    s['delta_potential_eur']=s.potential_eur-s.potential_v1_eur
    # Approved economic reference stays frozen: this change restores rural neutrality,
    # rather than silently retraining categories for unaffected metropolitan points.
    cut_source=ROOT/'pharmacy/output/archive/2027.2/thresholds.json'
    prior_manifest=json.loads((cut_source.parent/'manifest.json').read_text(encoding='utf-8'))
    if sha256(cut_source)!=prior_manifest['files_sha256']['thresholds.json']:
        raise ValueError('La referencia congelada de cortes 2027.2 ha cambiado')
    inputhashes[str(cut_source.relative_to(ROOT))]=sha256(cut_source)
    cutrows=json.loads(cut_source.read_text(encoding='utf-8'))
    allcuts={r['scope']:{'cuts_ascending_eur':[r['C_from_eur'],r['B_from_eur'],r['A_from_eur']]} for r in cutrows}
    s['category_national']=assign(s.potential_eur,allcuts['ESPAÑA']['cuts_ascending_eur'])
    s['category_ccaa']=''
    for ccaa,g in s.groupby('ccaa',sort=True):
        s.loc[g.index,'category_ccaa']=assign(g.potential_eur,allcuts[ccaa]['cuts_ascending_eur'])
    cuts=pd.DataFrame(cutrows)
    s['dual_category_disagreement']=s.category_national.ne(s.category_ccaa)
    for prefix in ['national','ccaa']:
        used=np.tile(allcuts['ESPAÑA']['cuts_ascending_eur'],(len(s),1)) if prefix=='national' else np.array([allcuts[c]['cuts_ascending_eur'] for c in s.ccaa])
        s[f'near_{prefix}_cut_5pct']=(np.abs(s.potential_eur.to_numpy()[:,None]/used-1).min(axis=1)<=.05)
    points=s.reset_index()
    sb,priority=sb_priority(points)
    for field in ['sb_priority_national','sb_priority_reason','national_AB_n','national_AB_eur','national_AB_market_pct','visitable_national_AB_n','visitable_national_AB_eur','visitable_national_AB_market_pct']:
        points[field]=points.sb_code.map(sb.set_index('sb_code')[field])
    points['category']=points.category_national # compatibility only; UI explicitly chooses its reference
    points['category_base']=points.category_national
    points['sensitive']=points.sensitive_v1 # historical diagnostic; always labelled V1 in UI
    points['sensitivity_variation']=points.sensitivity_variation_v1
    points['model_version']=VERSION
    original=universe['crm'][universe['original_columns']].reset_index(drop=True)
    crm=universe['crm'][universe['original_columns']+['establishment_id']].copy()
    fields=[c for c in points if c not in {'postal_code','crm_ids','category','category_base','sensitive','sensitivity_variation'}]
    crm=crm.merge(points[fields],left_on='establishment_id',right_on='id',how='left',sort=False,validate='many_to_one')
    assert crm[universe['original_columns']].equals(original)
    crm['contabilizar_modelo']=~crm.establishment_id.duplicated()
    crm['mercado_sumable_crm_eur']=np.where(crm.contabilizar_modelo,crm.potential_eur,0)
    cents=np.rint(points.potential_eur*100).astype('int64')
    actual=pd.Series(cents.to_numpy(),index=points.sb_code).groupby(level=0).sum().sort_index()
    assert actual.equals(market['budgets'].sort_index())
    assert len(points)==22316 and len(crm)==22356 and points.id.is_unique
    assert np.rint(crm.mercado_sumable_crm_eur*100).sum()==cents.sum()
    assert (points.pop65>=points.pop75).all() and (points.pop75>=points.pop85).all()
    assert np.allclose(points[AGES].sum(axis=1),points.pop65)
    assert np.allclose(points[AGES[2:]].sum(axis=1),points.pop75)
    for _,g in points.groupby('ccaa'):
        assert g.groupby('potential_eur').category_ccaa.nunique().max()==1
    controls={'version':VERSION,'crm_rows':len(crm),'physical_unique':len(points),'sb_count':len(sb),
       'market_cents':int(cents.sum()),'all_location_sb_conserved_cents':True,'all_original_crm_fields_unchanged':True,
       'population_unchanged_by_store':True,'ages_nested':True,'same_euros_for_both_grades':True,
       'sb_protections_applied':False,'scope_counts':points.geographic_scope.value_counts().to_dict(),
       'rural_neutral_weights':True,'rural_pools':int(points.loc[points.geographic_scope.eq('Rural'),'rural_group_key'].nunique()),
       'visit_eligible_count':int(points.visit_eligible.sum()),'rural_visit_eligible_count':int((points.geographic_scope.eq('Rural') & points.visit_eligible).sum()),
       'cuts_reference':'2027.2 frozen approved economic cuts','national_counts':points.category_national.value_counts().to_dict(),
       'ccaa_counts':points.category_ccaa.value_counts().to_dict(),'priority_sb_count':priority['priority_sb_count'],
       'dual_disagreements':int(points.dual_category_disagreement.sum()),'provisional_locations':int(points.location_provisional.sum()),
       'frozen_spatial_sources_hashes_verified':True,'raw_RX_manufacturer_market_reconciled':True}
    json_file(OUT/'controls.json',controls)
    json_file(OUT/'priority_policy.json',priority)
    json_file(OUT/'thresholds.json',cutrows)
    # Model tables are the one authoritative input of Excel and visualization, with identical IDs.
    json_file(OUT/'stores.json',points.to_dict('records'))
    json_file(OUT/'sanibricks.json',sb.to_dict('records'))
    csv(OUT/'universo_crm.csv',crm)
    csv(OUT/'sanibricks.csv',sb)
    csv(OUT/'sb_prioritarios.csv',sb[sb.sb_priority_national].sort_values('national_AB_eur',ascending=False))
    csv(OUT/'cortes_categorias.csv',cuts)
    if not args.no_excel:
        print('Exportando Excel nacional V2 y vistas de verificación...',flush=True)
        frame,distribution=export_book(points,crm,cuts,sb,priority,controls,inputhashes)
        csv(OUT/'establecimientos.csv',frame)
        csv(OUT/'distribucion_categorias.csv',distribution)
    else:
        frame=points.drop(columns=['category','category_base','sensitive','sensitivity_variation']).copy()
        frame['crm_ids']=frame.crm_ids.map(lambda v:' | '.join(map(str,v)))
        csv(OUT/'establecimientos.csv',frame)
    manifest={'version':VERSION,'generated_at':datetime.now(timezone.utc).isoformat(),'market_period':'Julio 2025–junio 2026',
       'market_scope':'ESSITY + RESTO RX sell-in','market_eur':int(cents.sum())/100,'inputs_sha256':inputhashes,
       'files_sha256':{p.name:sha256(p) for p in OUT.iterdir() if p.is_file() and p.name not in {'manifest.json','final_validation.json'}},
       'national_thresholds':allcuts['ESPAÑA']['cuts_ascending_eur'],'priority_policy':priority,
       'archive':'pharmacy/output/archive/2027.1','previous_delivery':'pharmacy/output/archive/2027.2','cuts_reference_version':'2027.2','new_origin_flows':False,
       'method':'Presupuesto SB ubicación × peso aplicado; rural igual municipio–SB, no rural señal espacial; cortes económicos congelados 2027.2.'}
    json_file(OUT/'manifest.json',manifest)
    print(json.dumps({'output':str(OUT),'controls':controls,'priority':priority},ensure_ascii=False),flush=True)


if __name__=='__main__':main()
